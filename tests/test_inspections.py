from ui_navigation import set_month
from ui_navigation import navigate
import unittest,tempfile,sys
from pathlib import Path
from io import BytesIO
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import inspections as api

class Inspections(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=self.tmp.name
 def tearDown(self):self.tmp.cleanup()
 def content(self):
  w=openpyxl.Workbook();s=w.active;s.title='Base';s.append(api.HEADERS)
  for os,end,actual,city in [('001','02/09/2026 10:00','02/09/2026 10:00','AGUDOS'),('001','02/09/2026 10:00','02/09/2026 10:01','AGUDOS'),('003','02/09/2026 10:00','03/09/2026 10:00','BOCAINA')]:s.append([os,city,'Fechada','Fechada','01/09/2026 00:00',end,actual,'VISTORIA'])
  b=BytesIO();w.save(b);return b.getvalue()
 def test_boundary_status_and_missing(self):
  r={'Status da OS':' Fechada ','Status da Operação':'FECHADA','Prazo de Execução - Fim':'02/09/2026 10:00','Data Fim Execução':'02/09/2026 10:00'}
  self.assertEqual(api.classification(r),'Dentro do prazo');r['Data Fim Execução']='02/09/2026 10:01';self.assertEqual(api.classification(r),'Fora do prazo');r['Data Fim Execução']='inválida';self.assertEqual(api.classification(r),'Pendente de data');r['Status da OS']='Aberta';self.assertEqual(api.classification(r),'Não concluída')
 def test_import_aggregate_period_repeat_export_and_fine(self):
  content=self.content();self.assertEqual(api.sheets(content),['Base']);self.assertEqual(api.import_file(self.root,content,'a.xlsx','Base'),3)
  with self.assertRaises(ValueError):api.import_file(self.root,content,'copia.xlsx','Base')
  total,rows,pending=api.summary(self.root,'2026-09-01','2026-09-30');self.assertEqual((total['total'],total['inside'],total['outside'],pending),(3,1,2,0));self.assertEqual(sum(r['total'] for r in rows),3)
  self.assertEqual(api.summary(self.root,'2026-09-02','2026-09-02')[0]['total'],2)
  with self.assertRaises(ValueError):api.fine(self.root,'2026-09-01','2026-09-30',10005,'Doc',False)
  with self.assertRaises(ValueError):api.fine(self.root,'2026-09-01','2026-09-30',10050,'Doc',True,(0,0))
  event=api.fine(self.root,'2026-09-01','2026-09-30',10050,'Doc e responsável',True);self.assertEqual(event['penalty_cents'],101);self.assertEqual(event['net_cents'],9949)
  w=openpyxl.load_workbook(BytesIO(api.export(self.root,'2026-09-01','2026-09-30',True)),read_only=True);self.assertEqual(w['Fora do prazo 1'].max_row,None);self.assertEqual(len(list(w['Fora do prazo 1'].values)),3);w.close()
  with api.connect(self.root) as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM visits').fetchone()[0],3);self.assertEqual(c.execute('SELECT COUNT(*) FROM penalties').fetchone()[0],1)
 def test_exact_duplicates_rollback_and_pending_blocks_fine(self):
  api.import_file(self.root,self.content(),'a.xlsx','Base')
  w=openpyxl.load_workbook(BytesIO(self.content()));w.active.append(['004','AGUDOS','Fechada','Fechada','','','02/09/2026 10:00','VISTORIA']);b=BytesIO();w.save(b)
  with self.assertRaises(Exception):api.import_file(self.root,b.getvalue(),'changed.xlsx','Base')
  with api.connect(self.root) as c:
   self.assertEqual(c.execute('SELECT COUNT(*) FROM imports').fetchone()[0],1)
   c.execute("UPDATE visits SET status='Pendente de data' WHERE id=1")
  self.assertEqual(api.summary(self.root,'2026-09-01','2026-09-30')[2],1)
  with self.assertRaises(ValueError):api.fine(self.root,'2026-09-01','2026-09-30',10000,'Doc',True)
 def test_ui_no_automatic_penalty_or_export(self):
  import os
  from datetime import date
  from unittest.mock import patch
  from streamlit.testing.v1 import AppTest
  root=Path(self.root)/'vistorias'
  api.import_file(root,self.content(),'a.xlsx','Base')
  import storage
  with patch.dict(os.environ,{'APP_DATA_DIR':self.root}),patch.object(storage,'DATA',Path(self.root)):
   app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run()
   set_month(app,date(2026,9,1))
   with patch.object(api,'export',side_effect=AssertionError('Exportação automática indevida')):
    navigate(app,module='Vistorias')
    navigate(app,'Prazos e Meta Contratual')
    self.assertFalse(app.exception)
    self.assertTrue(any(m.value=='99%' for m in app.metric))
    self.assertFalse(any(b.label=='Aplicar multa' for b in app.button))
   with api.connect(root) as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM penalties').fetchone()[0],0)
 def test_exact_target_no_penalty(self):
  with api.connect(self.root) as c:
   for i in range(100):c.execute('INSERT INTO visits(import_id,line,day,status,fingerprint,raw,city) VALUES(?,?,?,?,?,?,?)',(1,i,'2026-09-01','Dentro do prazo' if i<99 else 'Fora do prazo',str(i),'{}','AGUDOS'))
  self.assertTrue(api.summary(self.root,'2026-09-01','2026-09-30')[0]['met'])
  with self.assertRaises(ValueError):api.fine(self.root,'2026-09-01','2026-09-30',10000,'Doc',True)
if __name__=='__main__':unittest.main()
