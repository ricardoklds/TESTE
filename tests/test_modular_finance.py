import json,os,sqlite3,tempfile,unittest,hashlib
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
import openpyxl
import inspections as operational
import inspection_billing as bill
import contract_finance as contract

class ModularFinanceTests(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.data=Path(self.tmp.name);self.root=self.data/'vistorias'
 def tearDown(self):self.tmp.cleanup()
 def load(self,rows=None):
  w=openpyxl.Workbook();s=w.active;s.title='Número OS';h=[*operational.HEADERS,'Contrato','Causa Resultado','Formulário de Baixa','Família'];s.append(h)
  rows=rows or [('1','VISTORIA CADASTRAL','01/09/2026 10:00','Fechada'),('2','EXAME PREDIAL ALTA DE CONSUMO COMP MD','02/09/2026 10:01','Fechada'),('3','VERIFICAR USO DE FONTE ALTERNATIVA','02/09/2026 10:00','Fechada')]
  for os,tss,dt,status in rows:s.append([os,'BOCAINA',status,'Fechada','01/09/2026 00:00','02/09/2026 10:00',dt,tss,'C1','EXECUTADO','FORM','VISTORIA'])
  b=BytesIO();w.save(b);operational.import_file(self.root,b.getvalue(),'orig.xlsx','Número OS');bill.sync(self.root);return b.getvalue()
 def geo(self):bill.save_geo(self.root,'2026-09','Total mensal',10,{},'Documento Sabesp')
 def test_original_classification_exclusions_competence_and_operational(self):
  self.load();p=bill.preview(self.root,'2026-09');self.assertEqual(p['inspection_cents'],3450);self.assertEqual(p['pending'],0)
  self.assertEqual(operational.summary(self.root,'2026-09-01','2026-09-30')[0]['outside'],1)
  self.assertEqual(bill.preview(self.root,'2026-10')['inspection_cents'],0)
  self.assertEqual(sum(r['quantity'] for r in p['groups']),3)
  x=bill.reconciliation(self.root,'2026-09',{'200014104':999999},True)
  self.assertEqual(bill.preview(self.root,'2026-09')['inspection_cents'],3450)
  w=openpyxl.load_workbook(BytesIO(x),read_only=True);self.assertIn('Apuração por TSS',w.sheetnames);w.close()
 def test_duplicate_inactive_unknown_pending_and_individual_decision(self):
  self.load([('1','VISTORIA CADASTRAL','01/09/2026 10:00','Fechada'),('1','VISTORIA CADASTRAL','02/09/2026 10:00','Fechada'),('2','VISTORIA LIGAÇÃO INATIVA SEM DÉBITO','02/09/2026 10:00','Fechada'),('3','SERVIÇO DESCONHECIDO','02/09/2026 10:00','Fechada')])
  self.assertEqual(bill.preview(self.root,'2026-09')['pending'],4)
  bill.review(self.root,1,'2026-09','Cobrar','200014104','Execução válida comprovada');bill.review(self.root,2,'2026-09','Excluir',None,'Revisita não faturável')
  self.assertEqual(bill.preview(self.root,'2026-09')['inspection_cents'],1200);self.assertEqual(bill.preview(self.root,'2026-09')['pending'],2)
  with self.assertRaises(ValueError):bill.detail_page(self.root,'2026-09',limit=101)
 def test_prices_effectivity_freeze_import_block_and_backup(self):
  content=self.load();self.geo();name=next(s['name'] for s in bill.service_list(self.root,'2026-09') if s['code']=='200014104')
  bill.save_service(self.root,'200014104',name,'2026-10',1300,True,'Reajuste')
  self.assertEqual(bill.preview(self.root,'2026-09')['inspection_cents'],3450)
  self.assertEqual(next(s['cents'] for s in bill.service_list(self.root,'2026-10') if s['code']=='200014104'),1300)
  p=bill.preview(self.root,'2026-09');bill.approve(self.root,'2026-09',p['fingerprint'],'Conciliação conferida',True)
  self.assertEqual(bill.financial(self.root,'2026-09')['inspection_cents'],3450)
  with self.assertRaises(ValueError):bill.save_service(self.root,'200014104',name,'2026-09',1500,True,'Retroativo')
  with self.assertRaises(ValueError):bill.save_rule(self.root,'TSS','','','2026-09','Excluir',None,'Retroativo')
  with self.assertRaises(ValueError):bill.save_geo(self.root,'2026-09','Total mensal',20,{},'Alteração')
  w=openpyxl.load_workbook(BytesIO(content));w.active['A2']='NOVO';b=BytesIO();w.save(b)
  with self.assertRaises(ValueError):operational.import_file(self.root,b.getvalue(),'outro.xlsx','Número OS')
  backups=list((self.root/'backups').glob('*.sqlite3'));self.assertGreaterEqual(len(backups),2)
  with sqlite3.connect(backups[0]) as c:self.assertEqual(c.execute('PRAGMA integrity_check').fetchone()[0],'ok')
 def test_geo_two_modes_reference_no_double_count_and_history(self):
  self.load();bill.save_geo(self.root,'2026-09','Por município',3,{'AGUDOS':1,'BOCAINA':2},'Doc')
  p=bill.preview(self.root,'2026-09');self.assertEqual(p['geo']['cents'],3600);self.assertTrue(p['geo']['ready'])
  bill.save_geo(self.root,'2026-09','Total mensal',4,{'AGUDOS':1,'BOCAINA':2},'Correção');self.assertFalse(bill.preview(self.root,'2026-09')['geo']['ready'])
  with self.assertRaises(ValueError):bill.save_geo(self.root,'2026-09','Por município',None,{'Bocaina':1,'BOCAÍNA':2},'Doc')
  with bill.connect(self.root) as c:self.assertEqual(c.execute("SELECT COUNT(*) FROM audit WHERE action='Geolocalização'").fetchone()[0],2)
 def test_ambiguous_rules_pending_inactive_service_and_new_service(self):
  self.load();bill.save_service(self.root,'COD-CONFIRMADO','Novo serviço confirmado','2026-09',500,True,'Doc')
  bill.save_rule(self.root,'VISTORIA CADASTRAL','EXECUTADO','','2026-09','Cobrar','200014104','Doc')
  bill.save_rule(self.root,'VISTORIA CADASTRAL','','FORM','2026-09','Cobrar','COD-CONFIRMADO','Doc')
  self.assertEqual(bill.preview(self.root,'2026-09')['pending'],1)
  bill.save_rule(self.root,'VISTORIA CADASTRAL','','FORM','2026-09','Excluir',None,'Empate com exclusão')
  self.assertEqual(bill.preview(self.root,'2026-09')['pending'],1)
  bill.save_rule(self.root,'VISTORIA CADASTRAL','EXECUTADO','FORM','2026-09','Cobrar','COD-CONFIRMADO','Específica')
  self.assertEqual(bill.preview(self.root,'2026-09')['inspection_cents'],2750)
  bill.save_service(self.root,'COD-CONFIRMADO','Novo serviço confirmado','2026-09',500,False,'Suspenso')
  self.assertEqual(bill.preview(self.root,'2026-09')['pending'],1)
 def test_consolidation_readonly_validation_and_deduction_unique(self):
  self.load();self.geo();p=bill.preview(self.root,'2026-09');bill.approve(self.root,'2026-09',p['fingerprint'],'Doc',True)
  self.assertFalse(contract.preview(self.data,'2026-09')['ready'])
  with self.assertRaises(ValueError):contract.apply_fine(self.data,'2026-09','x','Doc',True)
  # Fonte de leituras simulada somente nesta unidade: a integração real é exercitada abaixo.
  with patch.object(contract,'readings',return_value={'cents':100000,'closed':True,'revision':1}):
   p=contract.preview(self.data,'2026-09');self.assertEqual(p['gross_cents'],115450)
   contract.approve(self.data,'2026-09',bill.digest(p),'Fechamento aprovado',True)
  s=contract.state(self.data,'2026-09');snapshot={k:v for k,v in s.items() if k not in ('validated','deduction','net_cents')}
  with self.assertRaises(ValueError):contract.apply_fine(self.data,'2026-09',bill.digest(snapshot),'Doc',False)
  d=contract.apply_fine(self.data,'2026-09',bill.digest(snapshot),'Doc',True);self.assertEqual(d['cents'],1155)
  with self.assertRaises(ValueError):contract.apply_fine(self.data,'2026-09',bill.digest(snapshot),'Doc',True)
  self.assertEqual(contract.state(self.data,'2026-09')['net_cents'],114295)
 def test_real_readings_adapter_preserves_six_totals_and_bank(self):
  import storage,queries,holidays
  with patch.object(storage,'DATA',self.data):
   import core
   from test_independent_validation import IndependentTests
   rows=[IndependentTests.row(self,1),IndependentTests.row(self,2,**{'Faturas Impressas':'Não'}),IndependentTests.row(self,3,**{'Causas de Leitura/Não Leitura Atual':'1'}),IndependentTests.row(self,4,**{'Fatura Retida':'Sim'})]
   IndependentTests.load(self,rows)
   holidays.save('2026-09-07','Teste','Nacional')
   prices=storage.month_info('2026-09')['prices']
   official=queries.summary({**queries.overview('2026-09'),**holidays.overview('2026-09')},prices)
   expected=sum(r['total_cents'] for r in official)
   before=hashlib.sha256((self.data/'contrato.sqlite3').read_bytes()).hexdigest()
   value=contract.readings(self.data,'2026-09');self.assertEqual(value['cents'],expected);self.assertEqual(value['holiday_quantity'],2)
   self.assertEqual(hashlib.sha256((self.data/'contrato.sqlite3').read_bytes()).hexdigest(),before)
 def test_import_incremental_projection_and_stale_approval(self):
  self.load();self.geo();old=bill.preview(self.root,'2026-09');self.assertEqual(bill.sync(self.root),0)
  bill.save_rule(self.root,'VISTORIA CADASTRAL','','','2026-09','Excluir',None,'Não faturável')
  with self.assertRaises(ValueError):bill.approve(self.root,'2026-09',old['fingerprint'],'Doc',True)
  self.assertEqual(len(bill.detail_page(self.root,'2026-09',offset=1,limit=1)),1)
 def test_independent_excel_calculation_and_exported_catalog(self):
  import validar_vistorias
  content=self.load();original=self.data/'orig.xlsx';original.write_bytes(content)
  p=bill.preview(self.root,'2026-09')
  r=validar_vistorias.run([original],'2026-09',self.data/'independent',bill.export_catalog(self.root,'2026-09'),{'200014104':999999},progress=lambda _:None)
  self.assertEqual(r['total_cents'],p['inspection_cents'])
  self.assertEqual(r['quantities']['200014104'],1)
  self.assertEqual(original.read_bytes(),content)
  with self.assertRaises(ValueError):validar_vistorias.run([original],'2026-09',self.data/'independent',progress=lambda _:None)
 def test_all_modules_lazy_and_no_exports_on_open(self):
  from streamlit.testing.v1 import AppTest
  import storage
  with patch.dict(os.environ,{'APP_DATA_DIR':str(self.data)}),patch.object(storage,'DATA',self.data),patch.object(bill,'reconciliation',side_effect=AssertionError('Exportação automática')):
   app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run()
   app.sidebar.selectbox[0].set_value('Vistorias').run()
   for page in ['Painel de Vistorias','Importação de Vistorias','Prazos e Meta Contratual','Serviços e Tabela de Preços','Geolocalização','Faturamento de Vistorias','Relatórios de Vistorias']:
    app.sidebar.radio[0].set_value(page).run();self.assertFalse(app.exception,page)
    if page=='Serviços e Tabela de Preços':
     next(s for s in app.selectbox if s.label=='Serviço').set_value('200014104').run()
     self.assertEqual(next(t for t in app.text_input if t.label=='Preço unitário em reais').value,'12,00')
   for module in ['Painel geral','Fechamento geral']:
    app.sidebar.selectbox[0].set_value(module).run();self.assertFalse(app.exception,module)
if __name__=='__main__':unittest.main()
