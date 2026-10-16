from ui_navigation import set_month
import hashlib,json,os,sqlite3,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from contextlib import contextmanager
from datetime import date
from io import BytesIO
import openpyxl
import inspection_billing as bill
import contract_finance as contract
from contract_ui import simulate
from ui_navigation import navigate
import test_modular_finance as fixtures

class SimpleInterface(unittest.TestCase):
 setUp=fixtures.ModularFinanceTests.setUp
 tearDown=fixtures.ModularFinanceTests.tearDown
 load=fixtures.ModularFinanceTests.load
 geo=fixtures.ModularFinanceTests.geo
 @contextmanager
 def app(self):
  from streamlit.testing.v1 import AppTest
  import storage
  with patch.dict(os.environ,{'APP_DATA_DIR':str(self.data)}),patch.object(storage,'DATA',self.data):
   a=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run();set_month(a,date(2026,9,1)).run()
   self.assertFalse(a.exception);yield a
 def bank_hashes(self):return {str(p.relative_to(self.data)):hashlib.sha256(p.read_bytes()).hexdigest() for p in self.data.rglob('*.sqlite3')}
 def metric(self,a,label):return next(m.value for m in a.metric if m.label==label)
 def test_menu_structure_no_module_selector_and_all_routes(self):
  self.load()
  import navigation
  with self.app() as a:
   self.assertFalse(any(s.label=='Módulo' for s in a.sidebar.selectbox));self.assertEqual([e.label for e in a.sidebar.expander],['Leituras','Vistorias','Configurações'])
   for module,pages in navigation.PAGES.items():
    for page in pages:
     navigate(a,page,module);self.assertFalse(a.exception,(module,page))
   self.assertNotIn('Serviços e Tabela de Preços',navigation.PAGES['Vistorias'])
   self.assertTrue(any(b.label=='Configurações de Leituras' for b in a.sidebar.button))
 def test_simulation_exact_rounding_and_off(self):
  self.assertEqual(simulate(9950,True),{'gross_cents':9950,'deduction_cents':100,'net_cents':9850})
  self.assertEqual(simulate(9949,True)['deduction_cents'],99)
  self.assertEqual(simulate(9950,False)['net_cents'],9950);self.assertEqual(simulate(0,True)['net_cents'],0)
 def test_toggle_never_writes_or_recalculates_financial_results(self):
  self.load();self.geo()
  with self.app() as a:
   navigate(a,module='Fechamento geral');self.assertFalse(a.exception);before=self.bank_hashes()
   self.assertEqual(self.metric(a,'Fechamento bruto'),'R$ 154,50')
   with patch.object(contract,'preview',side_effect=AssertionError('Recálculo financeiro')),patch.object(contract,'apply_fine',side_effect=AssertionError('Multa definitiva')),patch.object(bill,'preview',side_effect=AssertionError('Consulta financeira pesada')):
    a.toggle[0].set_value(True).run();self.assertFalse(a.exception);self.assertEqual(self.metric(a,'Glosa simulada'),'R$ 1,55');self.assertEqual(self.metric(a,'Fechamento líquido simulado'),'R$ 152,95')
    a.toggle[0].set_value(False).run();self.assertFalse(a.exception);self.assertEqual(self.metric(a,'Glosa simulada'),'R$ 0,00');self.assertEqual(self.metric(a,'Fechamento líquido simulado'),'R$ 154,50')
   self.assertEqual(self.bank_hashes(),before)
   with contract.connect(self.data) as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM deductions').fetchone()[0],0)
 def test_existing_definitive_deduction_is_separate_and_preserved(self):
  self.load();self.geo();p=bill.preview(self.root,'2026-09');bill.approve(self.root,'2026-09',p['fingerprint'],'Doc',True)
  with patch.object(contract,'readings',return_value={'cents':100000,'closed':True,'revision':1}):
   p=contract.preview(self.data,'2026-09');contract.approve(self.data,'2026-09',bill.digest(p),'Doc',True)
  official=contract.state(self.data,'2026-09');snapshot={k:v for k,v in official.items() if k not in ('validated','deduction','net_cents')};contract.apply_fine(self.data,'2026-09',bill.digest(snapshot),'Doc',True);official=contract.state(self.data,'2026-09')
  with self.app() as a:
   navigate(a,module='Fechamento geral');before=self.bank_hashes();self.assertEqual(self.metric(a,'Fechamento líquido simulado'),'R$ 1.154,50')
   a.toggle[0].set_value(True).run();a.toggle[0].set_value(False).run();self.assertFalse(a.exception);self.assertEqual(self.bank_hashes(),before)
   self.assertEqual(contract.state(self.data,'2026-09'),official)
   self.assertIn('Registro definitivo de glosa · fechamento oficial',[e.label for e in a.expander])
 def test_geo_save_without_document_live_total_and_no_double_count(self):
  self.load()
  with self.app() as a:
   navigate(a,'Geolocalização');self.assertFalse(a.exception);self.assertFalse(a.checkbox);self.assertFalse(any('Documento' in t.label or 'justificativa' in t.label for t in a.text_input))
   q=next(n for n in a.number_input if n.label=='Quantidade');self.assertIsNone(q.value);self.assertEqual(self.metric(a,'Preço unitário vigente'),'R$ 12,00')
   q.set_value(7).run();self.assertEqual(self.metric(a,'Valor total'),'R$ 84,00');next(b for b in a.button if b.label=='Salvar').click().run();self.assertFalse(a.exception)
   with bill.connect(self.root) as c:
    row=c.execute('SELECT * FROM geo').fetchone();self.assertEqual(row['total'],7);self.assertTrue(row['evidence']);self.assertEqual(c.execute("SELECT COUNT(*) FROM audit WHERE action='Geolocalização'").fetchone()[0],1)
   self.assertEqual(bill.preview(self.root,'2026-09')['geo']['cents'],8400)
   next(b for b in a.button if b.label=='Salvar').click().run();self.assertEqual(bill.preview(self.root,'2026-09')['geo']['cents'],8400)
 def test_geo_secondary_city_scope_switch_and_history(self):
  self.load()
  with self.app() as a:
   navigate(a,'Geolocalização');next(r for r in a.radio if r.label=='Forma de lançamento').set_value('Por município').run()
   next(t for t in a.text_area if t.label.startswith('Municípios e')).set_value('BOCAINA;2\nMACATUBA;10').run();self.assertEqual(self.metric(a,'Valor total'),'R$ 24,00')
   next(b for b in a.button if b.label=='Salvar').click().run();self.assertFalse(a.exception);self.assertEqual(bill.preview(self.root,'2026-09')['geo']['quantity'],2)
   next(r for r in a.radio if r.label=='Forma de lançamento').set_value('Total mensal').run();next(n for n in a.number_input if n.label=='Quantidade').set_value(5).run();next(b for b in a.button if b.label=='Salvar').click().run();self.assertFalse(a.exception)
   p=bill.preview(self.root,'2026-09');self.assertEqual(p['geo']['quantity'],5);self.assertEqual(p['geo']['cents'],6000)
   with bill.connect(self.root) as c:
    events=[json.loads(r[0]) for r in c.execute("SELECT payload FROM audit WHERE action='Geolocalização' ORDER BY id")];self.assertEqual(len(events),2);self.assertEqual(json.loads(events[-1]['before']['cities']),{'BOCAINA':2,'MACATUBA':10})
 def test_geo_validations_prices_effectivity_and_approved_lock(self):
  self.load();bill.save_geo(self.root,'2026-09','Total mensal',2,{})
  name=next(s['name'] for s in bill.service_list(self.root,'2026-09') if s['code']=='geo');bill.save_service(self.root,'geo',name,'2026-10',1300,True,'Reajuste')
  with self.app() as a:
   navigate(a,'Geolocalização');self.assertEqual(self.metric(a,'Preço unitário vigente'),'R$ 12,00');set_month(a,date(2026,10,1)).run();self.assertEqual(self.metric(a,'Preço unitário vigente'),'R$ 13,00')
  for quantity in [-1,1.5,True]:
   with self.assertRaises(ValueError):bill.save_geo(self.root,'2026-09','Total mensal',quantity,{})
  with self.assertRaises(ValueError):bill.save_geo(self.root,'2026-09','Por município',None,{'Bocaina':1,'BOCAÍNA':2})
  p=bill.preview(self.root,'2026-09');bill.approve(self.root,'2026-09',p['fingerprint'],'Doc',True)
  with self.app() as a:
   navigate(a,'Geolocalização');self.assertTrue(next(b for b in a.button if b.label=='Salvar').disabled)
  with self.assertRaises(ValueError):bill.save_geo(self.root,'2026-09','Total mensal',3,{})
 def test_general_report_official_currency_does_not_include_simulation(self):
  self.load();self.geo()
  with self.app() as a:
   navigate(a,module='Fechamento geral');a.toggle[0].set_value(True).run();next(b for b in a.button if b.label=='Preparar relatório do contrato').click().run();self.assertFalse(a.exception)
   content=a.session_state['contract_export'][4];w=openpyxl.load_workbook(BytesIO(content));s=w['Contrato'];self.assertEqual(s.cell(2,6).value,154.5);self.assertEqual(s.cell(2,7).value,0);self.assertEqual(s.cell(2,8).value,154.5);self.assertIn('R$',s.cell(2,6).number_format);self.assertIn('R$',s.cell(1,6).value);w.close()
if __name__=='__main__':unittest.main()
