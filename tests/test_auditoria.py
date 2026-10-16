"""Diagnóstico em bancos temporários; não altera os seis resultados financeiros."""
from contextlib import closing
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import openpyxl
import core,storage,holidays,auditoria_fechamento as audit,excel_io


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.previous=storage.DATA
        storage.DATA=Path(self.temp.name)/'data'
        with closing(storage.connect()):pass
        holidays.save('2026-09-07','Feriado de teste','Nacional')

    def tearDown(self):storage.DATA=self.previous;self.temp.cleanup()

    def import_fixture(self,name,values):
        book=openpyxl.Workbook();sheet=book.active;sheet.append(list(core.FIELDS.values()))
        for n,value in enumerate(values):
            row=['Sim','Não',None,2,'0','Funcionário teste','ATC AGUDOS',value,str(n),'Leitura','01']
            sheet.append(row)
        content=BytesIO();book.save(content)
        parsed=excel_io.inspect(content.getvalue())[0]
        return storage.import_sheet(content.getvalue(),name,{**parsed,'records':excel_io.iter_records(content.getvalue(),parsed)},parsed['mapping'],'2026-09')

    def test_input_comparison_arithmetic(self):
        rows=audit.comparison()
        self.assertEqual(sum(r['Diferença informada (centavos)'] for r in rows),140098)
        self.assertEqual(sum(audit.MANUAL[k]*audit.PRICES[k] for k in audit.KEYS),59691043)
        self.assertEqual(sum(audit.REPORTED[k]*audit.PRICES[k] for k in audit.KEYS),59550945)

    def test_date_classes_original_duplicates_financial_and_read_only(self):
        self.assertEqual(self.import_fixture('original.xlsx',['2026-09-07',None,'texto inválido','2026-10-01',44500,'2026-09-07T10:00:00.123']), (6,0))
        # Same raw rows in another file: change sheet title to alter file SHA only.
        original=next((storage.DATA/'originais').glob('*.xlsx'))
        book=openpyxl.load_workbook(original);book.active.title='Outra aba';buffer=BytesIO();book.save(buffer);book.close()
        p=excel_io.inspect(buffer.getvalue())[0]
        self.assertEqual(storage.import_sheet(buffer.getvalue(),'repetido.xlsx',{**p,'records':excel_io.iter_records(buffer.getvalue(),p)},p['mapping'],'2026-09'),(0,6))
        before=(storage.DATA/'contrato.sqlite3').read_bytes()
        original_bytes={p.name:p.read_bytes() for p in (storage.DATA/'originais').iterdir()}
        result=audit.audit(storage.DATA,Path(self.temp.name)/'resultado')
        self.assertTrue(result['available']);self.assertEqual(result['all_records'],6)
        self.assertEqual(result['integrity_issues'],{})
        self.assertEqual(result['all_missing_projected_date'],3)
        self.assertEqual(result['date_classes']['Data vazia'],1)
        self.assertEqual(result['date_classes']['Data fora do mês de referência'],1)
        self.assertEqual(result['date_classes']['Data interpretável; formato rejeitado pelo motor'],1)
        self.assertEqual(result['official_quantities']['138'],2)
        self.assertEqual(result['official_quantities']['impressao'],4)
        self.assertEqual(result['official_quantities']['apontamento'],2)
        self.assertEqual(result['official_quantities']['feriado'],2) # Fractional ISO accepted by holiday date projection independently.
        self.assertTrue(result['holiday_quantity_matches_sql'])
        self.assertEqual(result['imports'][1]['source_classifications']['Mesmo fingerprint ativo em outro arquivo (estado atual)'],6)
        self.assertEqual((storage.DATA/'contrato.sqlite3').read_bytes(),before)
        self.assertEqual({p.name:p.read_bytes() for p in (storage.DATA/'originais').iterdir()},original_bytes)
        with closing(openpyxl.load_workbook(Path(self.temp.name)/'resultado'/'auditoria.xlsx',read_only=True)) as x:
            self.assertEqual(len(list(x['Registros 1'].values)),7)
        with self.assertRaises(FileExistsError):audit.audit(storage.DATA,Path(self.temp.name)/'resultado')

    def test_absent_month_explicit_and_output_cannot_enter_data(self):
        before=(storage.DATA/'contrato.sqlite3').read_bytes()
        result=audit.audit(storage.DATA,Path(self.temp.name)/'vazio')
        self.assertFalse(result['available']);self.assertNotIn('official_quantities',result)
        self.assertEqual((storage.DATA/'contrato.sqlite3').read_bytes(),before)
        with self.assertRaises(ValueError):audit.audit(storage.DATA,storage.DATA/'auditoria')

    def test_excluded_and_closed_snapshots_audited_without_modification(self):
        self.import_fixture('arquivo.xlsx',['2026-09-07'])
        storage.set_import_deleted('2026-09',1,True,'Teste')
        result=audit.audit(storage.DATA,Path(self.temp.name)/'excluido',check_originals=False)
        self.assertEqual(result['official_quantities']['138'],0)
        self.assertEqual(result['official_quantities']['feriado'],0)
        storage.set_import_deleted('2026-09',1,False,'Teste')
        storage.close_month('2026-09')
        before=(storage.DATA/'contrato.sqlite3').read_bytes()
        result=audit.audit(storage.DATA,Path(self.temp.name)/'fechado',check_originals=False)
        self.assertTrue(result['closed']);self.assertEqual(result['official_quantities']['feriado'],1)
        self.assertEqual((storage.DATA/'contrato.sqlite3').read_bytes(),before)

    def test_orphan_projection_reported_without_changing_official_sum(self):
        self.import_fixture('arquivo.xlsx',['2026-09-07'])
        with closing(storage.connect()) as c:
            c.execute("INSERT INTO record_metrics VALUES(999,'2026-09',1,'2026-09-07','TESTE','ATC AGUDOS','01','GHOST',0,0,7,0,0,0,0,0,0)");c.commit()
        before=(storage.DATA/'contrato.sqlite3').read_bytes()
        result=audit.audit(storage.DATA,Path(self.temp.name)/'divergencia',check_originals=False)
        self.assertEqual(result['official_quantities']['138'],8)
        self.assertEqual(result['detail_quantities']['138'],1)
        self.assertFalse(result['detail_quantities_match_sql'])
        self.assertEqual(result['integrity_issues']['Projeções SQL sem registro original'],1)
        self.assertEqual((storage.DATA/'contrato.sqlite3').read_bytes(),before)
