"""Recálculo direto do Excel: não abre banco de produção nem lê a tabela dinâmica."""
import argparse,collections,datetime,hashlib,json,sqlite3,tempfile,unicodedata
from pathlib import Path
import openpyxl
from inspection_catalog import MAPPING,EXCLUDED
from inspection_contract import ALLOWED,EXCLUDED as OUTSIDE,MODES

def normalize(v):return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(v or '')) if not unicodedata.combining(c)).upper().split())
def initial_catalog():
    rules=[{'tss':normalize(tss),'cause':'','form':'','code':code,'decision':'Cobrar'} for code,items in MAPPING.items() for tss in items]
    rules.extend({'tss':normalize(tss),'cause':'','form':'','code':None,'decision':'Excluir'} for tss in EXCLUDED)
    return {'policy':{'allowed':ALLOWED,'excluded':OUTSIDE,'modes':MODES,'legacy':False},'prices':{'200014104':{'cents':1200,'active':1},'200014105':{'cents':2250,'active':1},'200014106':{'cents':2250,'active':1}},'rules':rules,'reviews':{},'scope':'Catálogo inicial confirmado em 15/10/2026, para setembro/2026 em diante.'}

def run(files,month,output,catalog=None,references=None,progress=print):
    """Algoritmo próprio: Excel em lotes + regras em Python + desempate em SQLite temporário."""
    datetime.date.fromisoformat(month+'-01');output=Path(output);output.mkdir(parents=True,exist_ok=True)
    if any((output/n).exists() for n in ['calculo_vistorias.json','conferencia_vistorias.xlsx']):raise ValueError('Saída já existe; escolha outra pasta para preservar os resultados anteriores.')
    if catalog is None and month<'2026-09':raise ValueError('Forneça catálogo com vigência comprovada para competência anterior a setembro/2026.')
    config=catalog or initial_catalog()
    if config.get('month',month)!=month:raise ValueError('O catálogo exportado pertence a outra competência.')
    periods=config.get('period_catalogs',{})
    result={'month':month,'files':[],'representatives':[],'representative_count':0,'scope':config.get('scope','Configuração exportada do sistema')}
    configurations={}
    def context(period):
        if period not in configurations:
            if period==month:cfg=config
            elif period in periods:cfg=periods[period]
            elif catalog is None and period>='2026-09':cfg=initial_catalog()
            else:raise ValueError('Exporte o catálogo de todos os períodos presentes nos arquivos: '+period)
            policies=cfg.get('policy',initial_catalog()['policy']);rules=collections.defaultdict(list)
            for rule in cfg['rules']:rules[normalize(rule['tss'])].append(rule)
            configurations[period]=(cfg,policies,rules)
        return configurations[period]
    headings=['Arquivo','Aba','Linha','OS','Execução','Município original','TSS original','Código','Motivo','Valor centavos','Competência','Ocorrência representante','Arquivo representante','Linha representante','Município representante','Execução representante']
    w=openpyxl.Workbook(write_only=True);sheet=w.create_sheet('Registros 1');sheet.append(headings)
    with tempfile.TemporaryDirectory(prefix='vistorias-independente-') as temp:
        with sqlite3.connect(Path(temp)/'scratch.sqlite3') as c:
            c.row_factory=sqlite3.Row
            c.execute('CREATE TABLE source(id INTEGER PRIMARY KEY,file TEXT,line INTEGER,os TEXT,contract TEXT,day TEXT,stamp TEXT,month TEXT,city TEXT,city_key TEXT,tss TEXT,tss_original TEXT,cause TEXT,form TEXT,closed INTEGER,raw_hash TEXT)')
            seen=set()
            ordered=list(map(Path,files));import_order={sha:i for i,sha in enumerate(config.get('import_order',[]))}
            if import_order:ordered.sort(key=lambda f:import_order.get(hashlib.sha256(f.read_bytes()).hexdigest(),len(import_order)))
            for file in ordered:
                before=hashlib.sha256(file.read_bytes()).hexdigest()
                if before in seen:raise ValueError('Mesmo arquivo selecionado duas vezes; nenhuma contagem será presumida.')
                seen.add(before);book=openpyxl.load_workbook(file,read_only=True,data_only=True)
                try:
                    rows=book['Número OS'].iter_rows(values_only=True);headers=next(rows)
                    if not all(h in headers for h in ['Número OS','Município','Status da OS','Status da Operação','Data Fim Execução','TSS']):raise ValueError('Cabeçalhos obrigatórios ausentes.')
                    batch=[];n=0
                    for line,values in enumerate(rows,2):
                        if not any(v is not None and str(v).strip() for v in values):continue
                        r=dict(zip(headers,values));value=r.get('Data Fim Execução')
                        try:
                            dt=value if isinstance(value,datetime.datetime) else datetime.datetime.strptime(str(value).strip(),'%d/%m/%Y %H:%M')
                            day=dt.date().isoformat();stamp=dt.isoformat(timespec='seconds')
                        except (ValueError,TypeError):day=stamp=None
                        raw_hash=hashlib.sha256(json.dumps(r,ensure_ascii=False,default=str).encode()).hexdigest();city=str(r.get('Município') or 'Não informado')
                        batch.append((file.name,line,str(r.get('Número OS') or ''),normalize(r.get('Contrato')),day,stamp,day[:7] if day else None,city,normalize(city),normalize(r.get('TSS')),str(r.get('TSS') or ''),normalize(r.get('Causa Resultado')),normalize(r.get('Formulário de Baixa')),int(all(normalize(r.get(k))=='FECHADA' for k in ('Status da OS','Status da Operação'))),raw_hash));n+=1
                        if len(batch)>=2000:c.executemany('INSERT INTO source VALUES(NULL,'+','.join('?'*15)+')',batch);batch=[];progress(f'{file.name}: {n} linhas lidas')
                    if batch:c.executemany('INSERT INTO source VALUES(NULL,'+','.join('?'*15)+')',batch)
                finally:book.close()
                if hashlib.sha256(file.read_bytes()).hexdigest()!=before:raise ValueError('Arquivo mudou durante a leitura.')
                result['files'].append({'name':file.name,'sha256':before,'rows':n})
            c.execute('CREATE TABLE decisions(id INTEGER PRIMARY KEY,code TEXT,cents INTEGER,mode TEXT,scope TEXT,reviewed INTEGER,reason TEXT)')
            frozen={(r['contract'],r['code'],r['os'],r['month']) for r in config.get('frozen_charges',[])}
            frozen_by_os=collections.defaultdict(set)
            for contract,code,os,period in frozen:frozen_by_os[(contract,code,os)].add(period)
            batch=[]
            for r in c.execute('SELECT * FROM source ORDER BY id'):
                cfg,pol,rules=context(r['month'] or month)
                candidates=[rule for rule in rules[r['tss']] if (not rule.get('cause') or normalize(rule['cause'])==r['cause']) and (not rule.get('form') or normalize(rule['form'])==r['form'])]
                if candidates:
                    score=lambda rule:bool(rule.get('cause'))+bool(rule.get('form'))
                    best=max(map(score,candidates));candidates=[rule for rule in candidates if score(rule)==best]
                review=cfg.get('reviews',{}).get(r['raw_hash']);rule=review or (candidates[0] if len(candidates)==1 else None);code=rule.get('code') if rule else None
                scope='Dentro do escopo' if pol.get('legacy') or r['city_key'] in pol['allowed'] else ('Fora do escopo do contrato' if r['city_key'] in pol['excluded'] else 'Pendente: município não autorizado')
                mode=pol['modes'].get(code,'manual');price=cfg['prices'].get(code)
                if scope=='Fora do escopo do contrato':reason=scope
                elif not r['closed']:reason='Não concluída'
                elif scope!='Dentro do escopo':reason=scope
                elif rule and rule['decision']=='Excluir':reason='Excluído por classificação confirmada'
                elif not review and r['tss']=='VISTORIA LIGACAO INATIVA SEM DEBITO':reason='Pendente: ligação inativa'
                elif not review and len(candidates)>1:reason='Pendente: regras ambíguas'
                elif rule and rule['decision']=='Pendente':reason='Pendente: classificação definida pelo usuário'
                elif not r['stamp']:reason='Pendente: data de execução sem horário válido'
                elif not r['os']:reason='Pendente: OS ausente'
                elif not rule:reason='Pendente: serviço não mapeado'
                elif not price or not price['active']:reason='Pendente: preço ausente ou serviço inativo'
                elif mode=='once' and frozen_by_os[(r['contract'],code,r['os'])]-{r['month']}:reason='Pendente: OS já faturada em competência aprovada'
                else:reason='Elegível'
                override=cfg.get('tss_prices',{}).get(r['tss'],{}).get('cents');cents=override if override is not None else (price['cents'] if price else 0)
                batch.append((r['id'],code,cents,mode,scope,int(review is not None),reason))
                if len(batch)>=2000:c.executemany('INSERT INTO decisions VALUES(?,?,?,?,?,?,?)',batch);batch=[]
            if batch:c.executemany('INSERT INTO decisions VALUES(?,?,?,?,?,?,?)',batch)
            # Uma tabela reduzida somente com candidatos; cada algoritmo decide sua elegibilidade antes de agrupar.
            c.execute("CREATE TEMP TABLE candidates AS SELECT s.id,s.contract,d.code,s.os,s.stamp,s.city_key FROM source s JOIN decisions d USING(id) WHERE d.reason='Elegível' AND d.mode='once'")
            c.execute('CREATE INDEX candidate_key ON candidates(contract,code,os,stamp DESC,id)')
            c.execute('CREATE TEMP TABLE winners AS SELECT contract,code,os,MIN(city_key) city_min,MAX(city_key) city_max,COUNT(*) n FROM candidates GROUP BY contract,code,os')
            c.execute('ALTER TABLE winners ADD COLUMN winner INTEGER')
            c.execute('UPDATE winners SET winner=(SELECT id FROM candidates x WHERE x.contract=winners.contract AND x.code=winners.code AND x.os=winners.os ORDER BY stamp DESC,id LIMIT 1)')
            c.execute('CREATE INDEX winner_key ON winners(contract,code,os)')
            c.execute("CREATE TEMP TABLE repetitions AS SELECT s.month,s.contract,d.code,s.os,SUM(d.scope='Dentro do escopo' AND s.closed) n,COUNT(*) legacy_n FROM source s JOIN decisions d USING(id) GROUP BY s.month,s.contract,d.code,s.os")
            c.execute('CREATE INDEX repetition_key ON repetitions(month,contract,code,os)')
            c.execute("CREATE TEMP TABLE executions AS SELECT s.contract,d.code,s.os,s.stamp,s.city_key,COUNT(*) n FROM source s JOIN decisions d USING(id) WHERE d.scope='Dentro do escopo' AND s.closed AND d.reason NOT IN ('Excluído por classificação confirmada','Pendente: classificação definida pelo usuário') GROUP BY s.contract,d.code,s.os,s.stamp,s.city_key")
            c.execute('CREATE INDEX execution_key ON executions(contract,code,os,stamp,city_key)')
            query="""SELECT s.*,d.code,d.cents,d.mode,d.reason,d.reviewed,w.city_min,w.city_max,w.winner,w.n candidate_n,r.n repeated,rr.file winner_file,rr.line winner_line,rr.city winner_city,rr.stamp winner_stamp,e.n execution_n
                FROM source s JOIN decisions d USING(id)
                LEFT JOIN winners w ON w.contract=s.contract AND w.code=d.code AND w.os=s.os
                LEFT JOIN source rr ON rr.id=w.winner
                LEFT JOIN repetitions r ON r.month=s.month AND r.contract=s.contract AND r.code=d.code AND r.os=s.os
                LEFT JOIN executions e ON e.contract=s.contract AND e.code=d.code AND e.os=s.os AND e.stamp=s.stamp AND e.city_key=s.city_key
                WHERE s.month=? OR (s.day IS NULL AND s.closed) ORDER BY s.id"""
            quantities=collections.Counter();amounts=collections.Counter();reasons=collections.Counter();tsstotals=collections.Counter();linecount=0
            for r in c.execute(query,(month,)):
                reason=r['reason']
                if reason=='Elegível':
                    if r['mode']=='once' and r['city_min']!=r['city_max']:reason='Pendente: OS com municípios divergentes'
                    elif r['mode']=='once' and r['id']!=r['winner']:reason='Sem cobrança adicional — OS já representada'
                    elif r['mode']=='execution' and (r['execution_n'] or 0)>1 and not r['reviewed']:reason='Pendente: execução sem identidade unívoca'
                    elif r['mode']=='manual' and (r['repeated'] or 0)>1 and not r['reviewed']:reason='Pendente: OS repetida'
                    else:reason='Faturável'
                cents=r['cents'] if reason=='Faturável' else 0
                if reason=='Faturável':quantities[r['code']]+=1;amounts[r['code']]+=cents
                if reason=='Faturável' and (r['candidate_n'] or 0)>1:
                    result['representative_count']+=1
                    if len(result['representatives'])<1000:result['representatives'].append({'os':r['os'],'file':r['file'],'line':r['line'],'municipality':r['city'],'month':r['month'],'executed_at':r['stamp'],'cents':cents})
                reasons[reason]+=1;tsstotals[(r['tss_original'],r['code'],reason)]+=1
                if linecount and linecount%1048575==0:sheet=w.create_sheet('Registros '+str(linecount//1048575+1));sheet.append(headings)
                from openpyxl.cell import WriteOnlyCell
                values=[r['file'],'Número OS',r['line'],r['os'],r['stamp'],r['city'],r['tss_original'],r['code'],reason,cents,r['month'],r['winner'],r['winner_file'],r['winner_line'],r['winner_city'],r['winner_stamp']];cells=[]
                for value in values:
                    if isinstance(value,str):cell=WriteOnlyCell(sheet,value=value);cell.data_type='s';cells.append(cell)
                    else:cells.append(value)
                sheet.append(cells);linecount+=1
                if linecount%2000==0:progress(f'{linecount} registros recalculados')
            result.update(quantities=dict(quantities),cents=dict(amounts),total_cents=sum(amounts.values()),reasons=dict(reasons),processed=linecount,catalog_sha256=hashlib.sha256(json.dumps(config,sort_keys=True,ensure_ascii=False).encode()).hexdigest())
            result['comparison']={code:{'calculated':quantities[code],'reference':ref,'difference':quantities[code]-ref} for code,ref in (references or {}).items()}
            s=w.create_sheet('Comparação');s.append(['Código','Quantidade calculada','Referência informada','Diferença'])
            for code,r in result['comparison'].items():s.append([code,r['calculated'],r['reference'],r['difference']])
            s=w.create_sheet('Resumo');s.append(['Código','Quantidade','Valor centavos'])
            for code in config['prices']:
                if code!='geo':s.append([code,quantities[code],amounts[code]])
            s=w.create_sheet('Motivos');s.append(['Motivo','Quantidade'])
            for key,n in reasons.items():s.append([key,n])
            s=w.create_sheet('TSS');s.append(['TSS original','Código','Motivo','Quantidade'])
            for (tss,code,reason),n in sorted(tsstotals.items(),key=lambda x:str(x[0])):s.append([tss,code,reason,n])
    (output/'calculo_vistorias.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');w.save(output/'conferencia_vistorias.xlsx');progress('Concluído. Nenhum banco de produção foi aberto.');return result
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('files',nargs='+',type=Path);p.add_argument('--mes',required=True);p.add_argument('--saida',type=Path,required=True);p.add_argument('--catalogo',type=Path);p.add_argument('--referencias',type=Path);a=p.parse_args()
    run(a.files,a.mes,a.saida,json.loads(a.catalogo.read_text(encoding='utf-8-sig')) if a.catalogo else None,json.loads(a.referencias.read_text(encoding='utf-8-sig')) if a.referencias else None)
