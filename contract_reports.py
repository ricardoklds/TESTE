"""Exportações gerenciais sob demanda. Recebe agregados, não abre bancos nem Excel de origem."""
from io import BytesIO
from xml.sax.saxutils import escape
from contract_presentation import money,unit_text,simulate

LABEL='PRÉVIA — NÃO APROVADA'

def financial_summary(composition,enabled,percent=1):
    rows=[('Leituras',composition['subtotals']['Leituras']),('Vistorias',composition['subtotals']['Vistorias']),('Geolocalização',composition['subtotals']['Geolocalização']),('Subtotal informativo: Vistorias + Geolocalização',composition['visits_geo_cents']),('Faturamento bruto',composition['gross_cents'])]
    if enabled:
        s=simulate(composition['gross_cents'],True,percent);label=f'Glosa simulada de {float(percent):g}%'.replace('.',',');rows.extend([(label,s['deduction_cents']),('Líquido simulado',s['net_cents'])])
    return rows

def excel(composition,enabled=False,percent=1):
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    w=openpyxl.Workbook();w.remove(w.active)
    def text(value):return ILLEGAL_CHARACTERS_RE.sub('',str(value))
    def append(s,values):
        s.append([text(v) if isinstance(v,str) else v for v in values])
        for cell in s[s.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    def currency(cell,cents):
        if cents is not None:cell.value=cents/100;cell.number_format='"R$" #,##0.00'
    s=w.create_sheet('Resumo Geral');append(s,['Gestão de Medição Sabesp']);append(s,[LABEL]);append(s,['Competência',composition['month']]);append(s,['Componente','Valor (R$)'])
    for label,amount in financial_summary(composition,enabled,percent):append(s,[label,None]);currency(s.cell(s.max_row,2),amount)
    append(s,['Observações'])
    for note in composition['notes']:append(s,[note])
    if enabled:append(s,['Glosa e líquido são somente simulações gerenciais, sem registro financeiro.'])
    for component,title in [('Leituras','Serviços de Leituras'),('Vistorias','Serviços de Vistorias'),('Geolocalização','Geolocalização')]:
        s=w.create_sheet(title);append(s,[LABEL]);append(s,['Competência',composition['month']]);append(s,['Código contratual','Descrição','Quantidade faturável','Valor unitário (R$)','Valor total (R$)','Detalhamento de preços'])
        for r in composition['rows']:
            if r['component']!=component:continue
            append(s,[r['code'] if r['code']!='geo' else 'Não informado (geo)',r['name'],r['quantity'],None,None,unit_text(r) if r.get('rate_details') else ''])
            currency(s.cell(s.max_row,4),r.get('unit_cents'));currency(s.cell(s.max_row,5),r['cents'])
        append(s,['Subtotal '+component,None,None,None,None]);currency(s.cell(s.max_row,5),composition['subtotals'][component])
        append(s,['Quantidades por serviço; não somar como fornecimentos distintos.'])
    for s in w:
        s.freeze_panes='A4';s.sheet_view.showGridLines=False
        for row in s:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for row in s:
            if row[0].value in ('Código contratual','Componente') or row[0].row==1:
                for cell in row:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163A60')
            if len(str(row[0].value or ''))>80:s.row_dimensions[row[0].row].height=15*(1+len(str(row[0].value))//75)
        for key,width in {'A':28,'B':60,'C':22,'D':23,'E':23,'F':48}.items():s.column_dimensions[key].width=width
        s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    out=BytesIO();w.save(out);return out.getvalue()

def pdf(composition,enabled=False,percent=1):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    out=BytesIO();doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=30)
    styles=getSampleStyleSheet();styles.add(ParagraphStyle('small_service',fontName='Helvetica',fontSize=8,leading=10));styles.add(ParagraphStyle('right_money',parent=styles['small_service'],alignment=TA_RIGHT))
    def p(text,style='small_service'):return Paragraph(escape(str(text)),styles[style])
    story=[p('Gestão de Medição Sabesp','Title'),p(LABEL,'Heading2'),p('Competência: '+composition['month']),Spacer(1,10)]
    for label,amount in financial_summary(composition,enabled,percent):story.append(p(label+': '+money(amount)))
    if enabled:story.append(p('Glosa e líquido são simulações gerenciais. Nenhum desconto foi registrado.'))
    story.append(Spacer(1,12))
    for component in composition['subtotals']:
        data=[[p(x) for x in ['Código contratual','Descrição','Quantidade faturável','Valor unitário','Valor total']]]
        for r in composition['rows']:
            if r['component']!=component:continue
            data.append([p(r['code'] if r['code']!='geo' else 'Não informado (geo)'),p(r['name']),p(r['quantity'] if r['quantity'] is not None else 'Não disponível','right_money'),p(unit_text(r),'right_money'),p(money(r['cents']),'right_money')])
        data.append([p('Subtotal'),p(component),p(''),p(''),p(money(composition['subtotals'][component]),'right_money')])
        table=Table(data,colWidths=[90,270,100,165,150],repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EFF7')),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#E7EFF7')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#163A60')),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]));story.extend([KeepTogether([p(component,'Heading2'),table]),Spacer(1,10)])
    story.append(p('Observações','Heading2'))
    story.extend(p(note) for note in composition['notes'])
    def page(canvas,doc):canvas.saveState();canvas.setFont('Helvetica',8);canvas.drawRightString(813,16,'Prévia gerencial · página '+str(doc.page));canvas.restoreState()
    doc.build(story,onFirstPage=page,onLaterPages=page);return out.getvalue()
