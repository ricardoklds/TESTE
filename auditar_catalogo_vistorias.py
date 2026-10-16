"""Investiga Excel original em bases temporárias. Não abre produção nem aprova faturamento."""
import argparse,collections,hashlib,json,tempfile
from pathlib import Path
import openpyxl
import inspections
import inspection_billing as billing
import validar_vistorias

TARGET='VERIFICAR SERVIÇO EXECUTADO'

def run(excel,month,output,references=None):
    excel=Path(excel);output=Path(output);output.mkdir(parents=True,exist_ok=False);original_hash=hashlib.sha256(excel.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix='auditoria-tss-') as temp:
        inspections.import_file(temp,excel.read_bytes(),excel.name,'Número OS');billing.sync(temp)
        # Reprodução explícita da configuração equivocada anterior; somente no banco temporário.
        billing.save_rule(temp,TARGET,'','',month,'Cobrar','200014104','Reprodução diagnóstica do mapeamento anterior')
        before=billing.preview(temp,month)
        change={'kind':'rule','tss':TARGET,'cause':'','form':'','effective':month,'decision':'Cobrar','code':'200014106','unit_cents':None,'evidence':'Correção comercial confirmada pelo usuário: Predial R$22,50'}
        impact=billing.prepare_change(temp,change);billing.confirm_change(temp,impact,True);after=billing.preview(temp,month)
        independent=validar_vistorias.run([excel],month,output/'independente',billing.export_catalog(temp,month),references=references,progress=lambda _:None)
        assert independent['total_cents']==after['inspection_cents']
        assert all(independent['quantities'].get(s['code'],0)==s['quantity'] for s in after['services'])
        (output/'conciliacao-catalogo.xlsx').write_bytes(billing.reconciliation(temp,month,references=references,detail=True))
        first,last=billing.month_bounds(month)
        with inspections.connect(temp) as c:
            rows=[dict(r) for r in c.execute('SELECT id,line,raw FROM visits WHERE day BETWEEN ? AND ? ORDER BY id',(first,last))]
        parsed=[(r['id'],r['line'],json.loads(r['raw'])) for r in rows]
        groups=collections.defaultdict(list)
        for id,line,r in parsed:groups[(billing.norm(r.get('Contrato')),str(r.get('Número OS') or ''))].append((id,line,r))
        repeated={k:v for k,v in groups.items() if k[1] and len(v)>1};pairs=[]
        for (_,os),rr in repeated.items():
            changed=[k for k in rr[0][2] if any(r.get(k)!=rr[0][2].get(k) for _,_,r in rr[1:])]
            if set(changed)<= {'Tempo Residual','Data Fim Execução'}:finding='Diferença somente no horário de conclusão/tempo residual; possível correção, não comprovada.'
            elif len({r.get('Causa Resultado') for _,_,r in rr})>1:finding='Resultados diferentes em ocorrências da mesma OS; não comprova duas cobranças.'
            else:finding='Mesmo resultado em datas distintas; confirmar regra comercial de revisita.'
            pairs.append({'os':os,'lines':[line for _,line,_ in rr],'changed_columns':changed,'finding':finding,'executions':[r.get('Data Fim Execução') for _,_,r in rr],'causes':[r.get('Causa Resultado') for _,_,r in rr]})
        target=[r for _,_,r in parsed if billing.norm(r.get('TSS'))==billing.norm(TARGET)]
        inactive=[r for _,_,r in parsed if billing.norm(r.get('TSS'))=='VISTORIA LIGACAO INATIVA SEM DEBITO']
        findings={'target_count':len(target),'target_causes':dict(collections.Counter(r.get('Causa Resultado') for r in target)),'target_status_os':dict(collections.Counter(r.get('Status da OS') for r in target)),'target_status_operation':dict(collections.Counter(r.get('Status da Operação') for r in target)),'duplicate_os':pairs,'duplicate_occurrences':sum(len(rr) for rr in repeated.values()),'inactive_count':len(inactive),'inactive_municipalities':dict(collections.Counter(r.get('Município') for r in inactive)),'inactive_causes':dict(collections.Counter(r.get('Causa Resultado') for r in inactive))}
        w=openpyxl.Workbook(write_only=True)
        headers=['ID temporário','Linha original','OS','Contrato','PDE','Município','TSS','Data competência','Fim execução','Prazo fim','Causa resultado','Formulário','Status OS','Status operação']
        def fields(id,line,r):return [id,line]+[r.get(k) for k in ['Número OS','Contrato','PDE','Município','TSS','Data de Competência','Data Fim Execução','Prazo de Execução - Fim','Causa Resultado','Formulário de Baixa','Status da OS','Status da Operação']]
        targets=w.create_sheet('TSS corrigido');targets.append(headers)
        pending=w.create_sheet('Ocorrências pendentes');pending.append(headers+['Motivo sem cobrança'])
        for id,line,r in parsed:
            if billing.norm(r.get('TSS'))==billing.norm(TARGET):targets.append(fields(id,line,r))
            key=(billing.norm(r.get('Contrato')),str(r.get('Número OS') or ''))
            if key in repeated:pending.append(fields(id,line,r)+['OS repetida: regra comercial não confirmada'])
            elif billing.norm(r.get('TSS'))=='VISTORIA LIGACAO INATIVA SEM DEBITO':pending.append(fields(id,line,r)+['Ligação inativa: sem enquadramento comercial confirmado'])
        s=w.create_sheet('Análise das OS');s.append(['OS','Linhas','Execuções','Causas','Colunas diferentes','Evidência/limitação'])
        for p in pairs:s.append([p['os'],str(p['lines']),' | '.join(p['executions']),' | '.join(p['causes']),', '.join(p['changed_columns']),p['finding']])
        w.save(output/'investigacao-registros.xlsx')
        result={'month':month,'original_sha256':original_hash,'before':before,'after':after,'impact':impact['impact'],'findings':findings,'comparison':independent['comparison'],'independent_quantities':independent['quantities'],'independent_sql_identical':True,'production_accessed':False,'billing_approved':False,'commercial_category_evidence':'Confirmação do usuário. Excel comprova descrição, registros e status, mas não contém preço/código tarifário.'}
        assert hashlib.sha256(excel.read_bytes()).hexdigest()==original_hash
        (output/'auditoria-catalogo.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));return result
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--excel',required=True,type=Path);p.add_argument('--mes',required=True);p.add_argument('--saida',required=True,type=Path);p.add_argument('--referencias',type=Path);a=p.parse_args();r=run(a.excel,a.mes,a.saida,json.loads(a.referencias.read_text()) if a.referencias else None);print(json.dumps({'before_cents':r['before']['inspection_cents'],'after_cents':r['after']['inspection_cents'],'pending':r['after']['pending']},indent=2))
