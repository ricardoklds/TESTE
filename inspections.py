"""Apuração operacional de prazo. Banco próprio; não modifica faturamento de leituras."""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
import sqlite3
import openpyxl
from io import BytesIO

class ClosingConnection(sqlite3.Connection):
    def __exit__(self,*args):
        try:return super().__exit__(*args)
        finally:self.close()

HEADERS=('Número OS','Município','Status da OS','Status da Operação','Prazo de Execução - Início','Prazo de Execução - Fim','Data Fim Execução','TSS')

def date_value(value):
    if isinstance(value,datetime):return value.isoformat(timespec='seconds')
    try:return datetime.strptime(str(value).strip(),'%d/%m/%Y %H:%M').isoformat()
    except (TypeError,ValueError):return None

def classification(raw):
    if any(str(raw.get(k,'')).strip().casefold()!='fechada' for k in ('Status da OS','Status da Operação')):return 'Não concluída'
    end=date_value(raw.get('Prazo de Execução - Fim')); executed=date_value(raw.get('Data Fim Execução'))
    if not end or not executed:return 'Pendente de data'
    return 'Dentro do prazo' if executed<=end else 'Fora do prazo'

def connect(root):
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    c=sqlite3.connect(root/'vistorias.sqlite3',factory=ClosingConnection);c.row_factory=sqlite3.Row
    c.execute('PRAGMA busy_timeout=30000')
    # A primeira versão cria somente banco novo. Bancos com versão desconhecida não são migrados.
    version=c.execute('PRAGMA user_version').fetchone()[0]
    exists=c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='visits'").fetchone()
    if exists and version!=1:c.close();raise ValueError('Versão desconhecida: nenhuma migração executada.')
    c.executescript('''CREATE TABLE IF NOT EXISTS imports(id INTEGER PRIMARY KEY,sha TEXT NOT NULL,sheet TEXT NOT NULL,name TEXT NOT NULL,created TEXT NOT NULL,UNIQUE(sha,sheet));
    CREATE TABLE IF NOT EXISTS visits(id INTEGER PRIMARY KEY,import_id INTEGER NOT NULL,line INTEGER NOT NULL,os TEXT,city TEXT,day TEXT,status TEXT NOT NULL,fingerprint TEXT NOT NULL UNIQUE,raw TEXT NOT NULL,UNIQUE(import_id,line));
    CREATE INDEX IF NOT EXISTS visit_period ON visits(day,status,city);
    CREATE TABLE IF NOT EXISTS penalties(id INTEGER PRIMARY KEY,created TEXT NOT NULL,payload TEXT NOT NULL);
    PRAGMA user_version=1;''')
    return c

def sheets(content):
    w=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        return [s.title for s in w if all(k in next(s.iter_rows(min_row=1,max_row=1,values_only=True),()) for k in HEADERS)]
    finally:w.close()

def import_file(root,content,name,sheet,progress=None):
    sha=hashlib.sha256(content).hexdigest();w=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    c=connect(root)
    try:
        # Apenas uma base operacional por arquivo: impede somar a cópia Planilha5.
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT 1 FROM imports WHERE sha=?',(sha,)).fetchone():raise ValueError('Este arquivo já possui uma aba importada. Nenhuma linha foi adicionada.')
        s=w[sheet];rows=s.iter_rows(values_only=True);headers=next(rows)
        if not all(k in headers for k in HEADERS):raise ValueError('Cabeçalhos obrigatórios ausentes.')
        dest=Path(root)/'originais';dest.mkdir(exist_ok=True);target=dest/(sha+'.xlsx')
        if target.exists():
            if hashlib.sha256(target.read_bytes()).hexdigest()!=sha:raise ValueError('Original arquivado inconsistente.')
        else:
            with target.open('xb') as f:f.write(content)
        import_id=c.execute('INSERT INTO imports(sha,sheet,name,created) VALUES(?,?,?,?)',(sha,sheet,name,datetime.now(timezone.utc).isoformat())).lastrowid
        batch=[];count=0
        for line,row in enumerate(rows,2):
            if not any(v is not None and str(v).strip() for v in row):continue
            raw=dict(zip(headers,row));dt=date_value(raw.get('Data Fim Execução'))
            city=str(raw.get('Município') or '').strip() or 'Não informado'
            batch.append((import_id,line,str(raw.get('Número OS') or ''),city,dt[:10] if dt else None,classification(raw),hashlib.sha256(json.dumps(raw,ensure_ascii=False,default=str,sort_keys=True).encode()).hexdigest(),json.dumps(raw,ensure_ascii=False,default=str)))
            count+=1
            if len(batch)>=2000:
                c.executemany('INSERT INTO visits(import_id,line,os,city,day,status,fingerprint,raw) VALUES(?,?,?,?,?,?,?,?)',batch);batch=[]
                if progress:progress(count)
        if batch:c.executemany('INSERT INTO visits(import_id,line,os,city,day,status,fingerprint,raw) VALUES(?,?,?,?,?,?,?,?)',batch)
        c.commit()
        if progress:progress(count)
        return count
    except Exception:c.rollback();raise
    finally:c.close();w.close()

def summary(root,start,end):
    with connect(root) as c:
        rows=[dict(r) for r in c.execute('''SELECT city,COUNT(*) total,
        SUM(status='Dentro do prazo') inside,SUM(status='Fora do prazo') outside
        FROM visits WHERE day BETWEEN ? AND ? AND status IN ('Dentro do prazo','Fora do prazo') GROUP BY city ORDER BY city''',(start,end))]
        pending=c.execute("SELECT COUNT(*) FROM visits WHERE status='Pendente de data' AND (day BETWEEN ? AND ? OR day IS NULL)",(start,end)).fetchone()[0]
    total={k:sum(r[k] for r in rows) for k in ('total','inside','outside')}
    def decorate(r):
        r['percent']=100*r['inside']/r['total'] if r['total'] else None
        r['outside_percent']=100*r['outside']/r['total'] if r['total'] else None
        r['met']=r['inside']*100>=r['total']*99 if r['total'] else None
        return r
    return decorate(total),[decorate(r) for r in rows],pending

def fine(root,start,end,base_cents,evidence,confirmed=False,expected_version=None):
    if not confirmed or not evidence.strip() or type(base_cents)!=int or base_cents<=0:raise ValueError('Confirmação, base positiva e documento do fechamento são obrigatórios.')
    # Ler índice e gravar evento na mesma transação; não altera fechamento oficial.
    c=connect(root)
    try:
        c.execute('BEGIN IMMEDIATE')
        version=tuple(c.execute('SELECT COUNT(*),COALESCE(MAX(id),0) FROM imports').fetchone())
        if expected_version is not None and tuple(expected_version)!=version:raise ValueError('Importações mudaram; atualize e confirme novamente.')
        counts=c.execute("SELECT COUNT(*),SUM(status='Dentro do prazo') FROM visits WHERE day BETWEEN ? AND ? AND status IN ('Dentro do prazo','Fora do prazo')",(start,end)).fetchone()
        pending=c.execute("SELECT COUNT(*) FROM visits WHERE status='Pendente de data' AND (day BETWEEN ? AND ? OR day IS NULL)",(start,end)).fetchone()[0]
        if pending or not counts[0] or counts[1]*100>=counts[0]*99:raise ValueError('Cálculo indisponível: meta atingida, período sem dados ou datas pendentes.')
        from decimal import Decimal,ROUND_HALF_UP
        penalty=int((Decimal(base_cents)/100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        event={'start':start,'end':end,'total':counts[0],'inside':counts[1],'base_cents':base_cents,'penalty_cents':penalty,'net_cents':base_cents-penalty,'evidence':evidence.strip(),'rate_percent':1,'target_percent':99,'imports':[dict(r) for r in c.execute('SELECT id,sha,sheet FROM imports ORDER BY id')]}
        c.execute('INSERT INTO penalties(created,payload) VALUES(?,?)',(datetime.now(timezone.utc).isoformat(),json.dumps(event,ensure_ascii=False)));c.commit();return event
    except Exception:c.rollback();raise
    finally:c.close()

def export(root,start,end,detail=False):
    w=openpyxl.Workbook(write_only=True);s=w.create_sheet('Resumo');total,rows,pending=summary(root,start,end)
    s.append(['Início','Fim','Total apurável','Dentro','Fora','Percentual dentro','Pendências (inclui sem data)']);s.append([start,end,total['total'],total['inside'],total['outside'],total['percent'],pending])
    s=w.create_sheet('Municípios');s.append(['Município','Total','Dentro','Fora','% dentro'])
    for r in rows:s.append([r['city'],r['total'],r['inside'],r['outside'],r['percent']])
    if detail:
        c=connect(root)
        try:
            n=0;page=0
            for r in c.execute("SELECT v.*,i.name,i.sheet FROM visits v JOIN imports i ON i.id=v.import_id WHERE day BETWEEN ? AND ? AND status='Fora do prazo' ORDER BY day,v.id",(start,end)):
                if n%1048575==0:
                    page+=1;s=w.create_sheet('Fora do prazo '+str(page));s.append(['Arquivo','Aba','Linha','OS','Município','TSS','Prazo início','Prazo fim','Tempo residual original','Execução','Situação'])
                raw=json.loads(r['raw']);s.append([r['name'],r['sheet'],r['line'],r['os'],r['city']]+[str(raw.get(k) or '') for k in ('TSS','Prazo de Execução - Início','Prazo de Execução - Fim','Tempo Residual','Data Fim Execução')]+[r['status']]);n+=1
        finally:c.close()
    b=BytesIO();w.save(b);return b.getvalue()
