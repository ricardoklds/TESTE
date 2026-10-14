"""Consultas de conferência somente leitura; não decide nem altera faturamento."""
from contextlib import closing
import json
from pathlib import Path
import sqlite3
import tempfile
import openpyxl
import holidays
from fast_reports import text_row

SITUATIONS=('Todos','Elegíveis cobrados','Pendências','Não cobrados')
PENDING=('Município/UF não confirmado','UF não confirmada','Execução no Brasil não comprovada',
         'Fornecimento/execução ausente','Domingo: conferir')
FIELDS={'id':'ID','day':'Data','municipality':'Município','state':'UF','city':'Cidade/ATC',
        'group_name':'Grupo','employee':'Funcionário','supply':'Fornecimento','execution':'Execução',
        'reason':'Motivo','country':'País','location_evidence':'Comprovação da localidade'}


def review_sql(c,month,start=None,end=None):
    frozen=bool(c.execute('SELECT 1 FROM holiday_freeze WHERE month=?',(month,)).fetchone())
    if frozen:
        base='SELECT '+','.join(f"json_extract(detail,'$.{k}') AS {k}" for k in FIELDS if k!='reason')
        base+=",json_extract(detail,'$.status') status FROM holiday_frozen_review WHERE month=?"
        args=[month]
        if start is not None:base+=" AND json_extract(detail,'$.day') BETWEEN ? AND ?";args.extend([start,end])
    else:base,args=holidays.candidate_sql(month,start,end)
    charged,charged_args=holidays.charges_query(c,month,start,end)
    if not frozen and base in charged and charged_args==args:
        # Reuse the exact candidates of the existing charge query, without repeating
        # its scan. Inner CTE names stay distinct from this outer materialization.
        charged=charged.replace(base,'SELECT * FROM review_candidates',1)
        charged_args=[]
    # Membership comes from the existing billing query, never from UI filters.
    pending=','.join('?' for _ in PENDING)
    divergent="c.status='Elegível' AND g.variants>1"
    sql='''WITH review_candidates AS MATERIALIZED ('''+base+'''),
        charged AS MATERIALIZED (SELECT id FROM ('''+charged+''')),
        executions AS (SELECT supply,day,execution,COUNT(DISTINCT municipality||'/'||state) variants
            FROM review_candidates WHERE status='Elegível' GROUP BY supply,day,execution)
        SELECT c.*,CASE WHEN h.id IS NOT NULL THEN 1 ELSE 0 END charged,
        CASE WHEN c.status IN ('''+pending+')'+('' if frozen else ' OR ('+divergent+')')+''' THEN 1 ELSE 0 END pending,
        CASE WHEN h.id IS NOT NULL THEN 'Elegível cobrado'
             WHEN '''+divergent+''' THEN 'Localidades divergentes na mesma execução'
             WHEN c.status='Elegível' THEN 'Execução duplicada: adicional já atribuído ou congelado sem cobrança'
             ELSE c.status END reason
        FROM review_candidates c LEFT JOIN charged h ON h.id=c.id
        LEFT JOIN executions g ON g.supply=c.supply AND g.day=c.day AND g.execution=c.execution
        WHERE c.status<>'Não aplicável' '''
    return sql,[*args,*charged_args,*PENDING]


def filtered_sql(c,month,start=None,end=None,municipality=None,group=None,employee=None,situation='Todos'):
    if situation not in SITUATIONS:raise ValueError('Situação inválida.')
    sql,args=review_sql(c,month,start,end)
    condition={'Todos':'1','Elegíveis cobrados':'charged=1','Pendências':'pending=1','Não cobrados':'charged=0'}[situation]
    for key,value in [('municipality',municipality),('group_name',group),('employee',employee)]:
        if value is not None:
            expression="COALESCE(municipality,'')||'/'||COALESCE(state,'')" if key=='municipality' else key
            condition+=' AND '+expression+'=?';args.append(value)
    return 'SELECT * FROM ('+sql+') WHERE '+condition,args


def display(row):
    result=dict(row)
    result.update(status=result['reason'],Cobrado=bool(result['charged']),**{'Valor (R$)':1.05 if result['charged'] else 0})
    return result


def page(month,start=None,end=None,offset=0,limit=100,municipality=None,group=None,employee=None,situation='Todos',known_count=None):
    if not 1<=limit<=500 or offset<0:raise ValueError('Página inválida.')
    with holidays.connection() as c:
        sql,args=filtered_sql(c,month,start,end,municipality,group,employee,situation)
        count=known_count if known_count is not None else c.execute('SELECT COUNT(*) FROM ('+sql+')',args).fetchone()[0]
        rows=c.execute(sql+' ORDER BY day,id LIMIT ? OFFSET ?',[*args,limit,offset]).fetchall()
    return count,[display(r) for r in rows]


def statistics(month,start=None,end=None,municipality=None,group=None,employee=None):
    """Poucos grupos de motivo, compartilháveis entre filtros e páginas."""
    with holidays.connection() as c:
        sql,args=filtered_sql(c,month,start,end,municipality,group,employee)
        return [dict(r) for r in c.execute('SELECT reason,charged,pending,COUNT(*) n FROM ('+sql+') GROUP BY reason,charged,pending ORDER BY n DESC,reason',args)]


def pending_summary(month,start=None,end=None,municipality=None,group=None,employee=None):
    with holidays.connection() as c:
        sql,args=filtered_sql(c,month,start,end,municipality,group,employee,'Pendências')
        return [dict(r) for r in c.execute('SELECT reason AS "Motivo",COUNT(*) AS "Quantidade" FROM ('+sql+') GROUP BY reason ORDER BY COUNT(*) DESC,reason',args)]


def pending_excel(db_path,revision,month,start=None,end=None,municipality=None,group=None,employee=None,output_dir=None):
    """Snapshot de leitura transacional, cursor SQL e XLSX incremental, em worker.

    Não copia/carrega o histórico inteiro. O arquivo final pode ocupar RAM ao baixar,
    conforme o Streamlit; só é preparado mediante solicitação explícita.
    """
    folder=Path(output_dir) if output_dir else Path(db_path).parent/'exports'
    folder.mkdir(parents=True,exist_ok=True)
    handle=tempfile.NamedTemporaryFile(prefix=f'pendencias-feriados-{month}-',suffix='.xlsx',dir=folder,delete=False)
    target=Path(handle.name);handle.close()
    try:
        with closing(sqlite3.connect(Path(db_path).resolve().as_uri()+'?mode=ro',uri=True)) as c:
            c.row_factory=sqlite3.Row;c.execute('BEGIN')
            if c.execute('SELECT revision FROM cache_revision WHERE id=1').fetchone()[0]!=revision:
                raise ValueError('A base mudou. Solicite a exportação novamente.')
            sql,args=filtered_sql(c,month,start,end,municipality,group,employee,'Pendências')
            workbook=openpyxl.Workbook(write_only=True)
            sheet=workbook.create_sheet('Pendências 1');text_row(sheet,list(FIELDS.values()))
            row_count=0;sheet_index=1
            for row in c.execute(sql+' ORDER BY day,id',args):
                if row_count==1048575:
                    sheet_index+=1;sheet=workbook.create_sheet(f'Pendências {sheet_index}')
                    text_row(sheet,list(FIELDS.values()));row_count=0
                text_row(sheet,[row[k] for k in FIELDS]);row_count+=1
            summary=workbook.create_sheet('Resumo por motivo');text_row(summary,['Motivo','Quantidade'])
            for row in c.execute('SELECT reason,COUNT(*) FROM ('+sql+') GROUP BY reason ORDER BY COUNT(*) DESC,reason',args):
                text_row(summary,list(row))
            metadata=workbook.create_sheet('Filtros')
            for key,value in [('Mês',month),('Início',start),('Fim',end),('Município/UF',municipality),('Grupo',group),('Funcionário',employee),('Situação','Pendências'),('Revisão',revision)]:text_row(metadata,[key,value])
            workbook.save(target)
        return target
    except Exception:
        target.unlink(missing_ok=True)
        raise
