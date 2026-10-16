"""Visões gerenciais somente leitura. Não altera cálculos nem seleciona novos feriados."""
from contextlib import closing
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import openpyxl
import holidays
from fast_reports import text_row

DIMS={'city':'city','group':'group_name','employee':'employee'}
LABELS={'city':'Cidade/ATC da origem','group':'Grupo','employee':'Funcionário'}
UNITS={'138':'u138','128':'u128','137':'u137','impressao':'uprint','apontamento':'uop','feriado':'fee'}
MEASURES=('records','undated','errors',*UNITS.values())


def connect(path):
    c=sqlite3.connect((Path(path).resolve()/'contrato.sqlite3').as_uri()+'?mode=ro',uri=True)
    c.row_factory=sqlite3.Row;c.execute('PRAGMA busy_timeout=30000');c.execute('BEGIN');return c


def dimension(alias,key):
    col=alias+'.'+DIMS[key]
    return "CASE WHEN NULLIF(TRIM("+col+"),'') IS NULL THEN 'Não informado' ELSE "+col+' END'


def conditions(alias,month,start,end,filters):
    parts=[alias+'.month=?'];args=[month]
    if start is not None:parts.append(alias+'.day BETWEEN ? AND ?');args.extend([start,end])
    for key in DIMS:
        values=tuple(sorted(set(filters.get(key,()))))
        if values:
            parts.append(dimension(alias,key)+' IN (SELECT value FROM json_each(?))');args.append(json.dumps(values,ensure_ascii=False))
    return ' AND '.join(parts),args


def aggregate_sql(c,month,start=None,end=None,filters=None,by=()):
    filters=filters or {}
    if any(k not in DIMS for k in by):raise ValueError('Dimensão inválida')
    where,args=conditions('m',month,start,end,filters)
    keys=','.join(dimension('m',k)+' AS "'+k+'"' for k in by);group=' GROUP BY '+','.join('"'+k+'"' for k in by) if by else ''
    # Processados inclui erros; unidades seguem exatamente os critérios das telas oficiais.
    valid='m.errors=0' if start is not None else '1'
    base='SELECT '+(keys+',' if keys else '')+'COUNT(*) records,COALESCE(SUM(m.day IS NULL),0) undated,COALESCE(SUM(m.errors>0),0) errors,'+','.join(
        'COALESCE(SUM(CASE WHEN '+valid+' THEN m.'+col+' ELSE 0 END),0) '+col for col in list(UNITS.values())[:-1])+',0 fee FROM record_metrics m WHERE '+where+' AND m.excluded=0 AND m.import_deleted=0'+group
    # Selecionar os vencedores globais ANTES dos filtros impede realocar cobranças duplicadas.
    fees,fee_args=holidays.charges_query(c,month,start,end)
    fwhere,fargs=conditions('h',month,start,end,filters)
    # charges_query não expõe month em todos os tipos de retorno; a competência já está na subconsulta.
    fwhere=fwhere.replace('h.month=?','1',1);fargs=fargs[1:]
    fkeys=','.join(dimension('h',k)+' AS "'+k+'"' for k in by)
    extras='SELECT '+(fkeys+',' if fkeys else '')+'0 records,0 undated,0 errors,0 u138,0 u128,0 u137,0 uprint,0 uop,COUNT(*) fee FROM ('+fees+') h WHERE '+fwhere+group
    fields=','.join('"'+k+'"' for k in by);totals=','.join('COALESCE(SUM('+v+'),0) '+v for v in MEASURES)
    sql='SELECT '+(fields+',' if fields else '')+totals+' FROM ('+base+' UNION ALL '+extras+')'+group
    return sql,[*args,*fee_args,*fargs]


def decorate(row,prices):
    d=dict(row);d['readings']=d['u138']+d['u128']+d['u137']
    d['normal_cents']=sum(d[UNITS[k]]*prices[k] for k in UNITS if k!='feriado')
    d['holiday_cents']=d['fee']*105;d['total_cents']=d['normal_cents']+d['holiday_cents'];return d


def overview(path,month,prices,start=None,end=None,filters=None):
    with closing(connect(path)) as c:
        sql,args=aggregate_sql(c,month,start,end,filters)
        return decorate(c.execute(sql,args).fetchone(),prices)


def options(path,month,start=None,end=None):
    with closing(connect(path)) as c:
        result={}
        # Uma lista por dimensão, sem carregar os registros nem ler Excel.
        for key in DIMS:
            where,args=conditions('m',month,start,end,{})
            result[key]=[r[0] for r in c.execute('SELECT DISTINCT '+dimension('m',key)+' name FROM record_metrics m WHERE '+where+' AND excluded=0 AND import_deleted=0 ORDER BY name',args)]
        return result


def page(path,month,prices,start=None,end=None,filters=None,by=('employee',),offset=0,limit=50):
    if not 1<=limit<=100 or offset<0:raise ValueError('Paginação inválida')
    with closing(connect(path)) as c:
        sql,args=aggregate_sql(c,month,start,end,filters,by)
        order=' ORDER BY u138+u128+u137 DESC,'+','.join('"'+k+'"' for k in by)
        selected=c.execute('SELECT *,COUNT(*) OVER() _groups FROM ('+sql+')'+order+' LIMIT ? OFFSET ?',[*args,limit,offset]).fetchall()
        n=selected[0]['_groups'] if selected else c.execute('SELECT COUNT(*) FROM ('+sql+')',args).fetchone()[0]
        rows=[decorate(r,prices) for r in selected]
        for row in rows:row.pop('_groups',None)
        return n,rows


def excel(path,month,prices,start,end,filters,revision):
    with closing(connect(path)) as c:
        if c.execute('SELECT revision FROM cache_revision WHERE id=1').fetchone()[0]!=revision:raise ValueError('A base mudou. Atualize a tela antes de exportar.')
        book=openpyxl.Workbook(write_only=True);s=book.create_sheet('Escopo gerencial')
        for pair in [('Tipo','Relatório gerencial — não altera o fechamento oficial'),('Competência',month),('Início',start or 'Mês completo'),('Fim',end or 'Mês completo'),('Filtros',json.dumps(filters,ensure_ascii=False)),('Revisão',revision),('Localidade','Cidade/ATC informada na origem; não pressupõe município confirmado')]:text_row(s,list(pair))
        sql,args=aggregate_sql(c,month,start,end,filters);total=decorate(c.execute(sql,args).fetchone(),prices)
        for key in ('normal_cents','holiday_cents','total_cents','records','readings','undated'):text_row(s,[key,total[key]])
        s=book.create_sheet('Seis serviços');text_row(s,['Serviço','Quantidade','Preço centavos','Total centavos'])
        for key,col in UNITS.items():
            price=105 if key=='feriado' else prices[key];text_row(s,[key,total[col],price,total[col]*price])
        for dim,title in [('employee','Funcionários'),('city','Cidades ATC'),('group','Grupos')]:
            s=book.create_sheet(title);text_row(s,[LABELS[dim],'Registros processados','Leituras faturáveis','Normais centavos','Feriados centavos','Total centavos'])
            sql,args=aggregate_sql(c,month,start,end,filters,(dim,))
            for row in c.execute(sql+' ORDER BY u138+u128+u137 DESC,"'+dim+'"',args):
                r=decorate(row,prices);text_row(s,[r[dim],r['records'],r['readings'],r['normal_cents'],r['holiday_cents'],r['total_cents']])
        # Agregados completos, sem limite de ranking; cursor incremental e divisão de abas.
        sql,args=aggregate_sql(c,month,start,end,filters,tuple(DIMS))
        sheet=None;n=0;index=0
        for row in c.execute(sql+' ORDER BY employee,city,"group"',args):
            if n%1048575==0:
                index+=1;sheet=book.create_sheet('Detalhamento '+str(index));text_row(sheet,['Cidade/ATC','Grupo','Funcionário','Registros processados','Leituras faturáveis','Normais centavos','Feriados centavos','Total centavos'])
            r=decorate(row,prices);text_row(sheet,[r['city'],r['group'],r['employee'],r['records'],r['readings'],r['normal_cents'],r['holiday_cents'],r['total_cents']]);n+=1
        out=BytesIO();book.save(out);return out.getvalue()


def pdf(data,month,start,end,filters,ranking,detail):
    """Resumo gerencial limitado e identificado; Excel contém todos os agregados."""
    from xml.sax.saxutils import escape
    from reportlab.lib.pagesizes import landscape,A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from reportlab.lib import colors
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from exports import money
    out=BytesIO();styles=getSampleStyleSheet();story=[]
    def text(value):return Paragraph(escape(str(value)),styles['Normal'])
    story.append(Paragraph('Relatório gerencial — Gestão de Medição',styles['Title']))
    story.append(text(f'{month} | {start or "Mês completo"} a {end or "Mês completo"}'))
    story.append(text('Filtros: '+json.dumps(filters,ensure_ascii=False)))
    story.append(text('Visão gerencial. Não altera nem substitui o fechamento oficial. Cidade/ATC conforme origem.'))
    for label,key in [('Leituras e adicionais normais','normal_cents'),('Adicional de feriados','holiday_cents'),('Faturamento total','total_cents')]:story.append(text(label+': '+money(data[key])))
    story.append(text(f'Registros processados: {data["records"]}; leituras faturáveis: {data["readings"]}; sem data: {data["undated"]}.'))
    chart=VerticalBarChart();chart.x=45;chart.y=25;chart.width=300;chart.height=100
    chart.data=[(data['normal_cents']/100,data['holiday_cents']/100)];chart.categoryAxis.categoryNames=['Normais','Feriados'];chart.valueAxis.valueMin=0
    if data['total_cents']==0:chart.valueAxis.valueMax=1
    drawing=Drawing(370,150);drawing.add(chart);story.extend([drawing,Spacer(1,12)])
    for title,rows,dims in [('Ranking: até 50 funcionários',ranking,('employee',)),('Detalhamento: até 100 combinações da primeira página',detail,('city','group','employee'))]:
        story.append(Paragraph(title,styles['Heading2']))
        values=[[text(LABELS[k]) for k in dims]+[text(v) for v in ['Processados','Leituras','Normais','Feriados','Total']]]
        for r in rows:values.append([text(r[k]) for k in dims]+[text(v) for v in [r['records'],r['readings'],money(r['normal_cents']),money(r['holiday_cents']),money(r['total_cents'])]])
        table=Table(values,repeatRows=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8f0f6')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.gray)]));story.extend([table,Spacer(1,12)])
    story.append(text('Excel gerencial contém todas as combinações da seleção. Nenhum funcionário ou município foi atribuído por suposição.'))
    SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=20,rightMargin=20).build(story)
    return out.getvalue()
