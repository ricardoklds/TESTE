from collections import defaultdict
from io import BytesIO
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from core import CATEGORIES, summarize

def money(cents):
    text = f'{cents / 100:,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')
    return f'R$ {text}'

def productivity(rows, prices, by):
    groups = defaultdict(lambda: {'Registros':0, 'Leituras faturáveis':0, 'Comunicados':0, 'Apontamentos':0, 'total_cents':0})
    for row in rows:
        if row['excluded']:
            continue
        result = row['result']; units = result['units']; g = groups[result[by]]
        g['Registros'] += 1
        g['Leituras faturáveis'] += sum(units[k] for k in ('138','128','137'))
        g['Comunicados'] += units['impressao']; g['Apontamentos'] += units['apontamento']
        g['total_cents'] += sum(units[k] * prices[k] for k in CATEGORIES)
    label = 'Funcionário' if by == 'employee' else 'Unidade / cidade (ATC)'
    return [{label: name, **v} for name,v in sorted(groups.items(), key=lambda x:-x[1]['Leituras faturáveis'])]

def excel_report(month, rows, prices, imports, audit, contract, closed, goals=None):
    goals = goals or {}
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Fechamento'
    sheet.append(['Contrato', contract.get('name','Não informado')]); sheet.append(['Número',contract.get('number','Não informado')])
    sheet.append(['Referência',month]); sheet.append(['Situação','Fechado' if closed else 'Prévia'])
    sheet.append(['Código','Categoria','Quantidade','Valor unitário (R$)','Total (R$)'])
    total = 0
    for r in summarize(rows,prices):
        sheet.append([r['Código'],r['Categoria'],r['Quantidade'],r['unit_cents']/100,r['total_cents']/100]); total += r['total_cents']
    sheet.append(['','TOTAL','','',total/100])
    for row in sheet.iter_rows(min_row=6):
        for cell in (row[3],row[4]):cell.number_format = '"R$" #,##0.00'
    for by, title in [('employee','Funcionários'),('city','Unidades')]:
        s = book.create_sheet(title); entries = productivity(rows,prices,by)
        s.append([('Funcionário' if by=='employee' else 'Unidade / cidade (ATC)'), 'Registros','Leituras faturáveis','Comunicados','Apontamentos','Total (R$)'])
        for r in entries:s.append([*list(r.values())[:-1],r['total_cents']/100])
        for row in s.iter_rows(min_row=2):row[-1].number_format='"R$" #,##0.00'
    s=book.create_sheet('Metas');s.append(['Dimensão','Funcionário / ATC','Leituras faturáveis','Meta mensal','Atingimento (%)'])
    for by in ('employee','city'):
        for r in productivity(rows,prices,by):
            name=next(iter(r.values()));goal=goals.get(by+'|'+name)
            s.append(['Funcionário' if by=='employee' else 'ATC',name,r['Leituras faturáveis'],goal,round(100*r['Leituras faturáveis']/goal,1) if goal else None])
    s = book.create_sheet('Conferência')
    s.append(['ID','Arquivo','Aba','Linha original','Excluído','Motivo','Erros','Avisos',*CATEGORIES.keys(),'Total (R$)'])
    for row in rows:
        u=row['result']['units']; total=sum(u[k]*prices[k] for k in CATEGORIES) if not row['excluded'] else 0
        s.append([row['id'],row['name'],row['sheet'],row['line'],bool(row['excluded']),row['reason'], '; '.join(row['result']['errors']),'; '.join(row['result']['alerts']),*[u[k] for k in CATEGORIES],total/100])
    raw = book.create_sheet('Dados importados')
    headers = list(dict.fromkeys(h for row in rows for h in row['raw']))
    raw.append(['ID do registro',*headers])
    for row in rows:raw.append([row['id'],*[row['raw'].get(h) for h in headers]])
    s=book.create_sheet('Arquivos e integridade');s.append(['Arquivo','Aba','SHA-256','Importado em UTC','Mapeamento'])
    for r in imports:s.append([r['name'],r['sheet'],r['sha'],r['created_at'],r['mapping']])
    s=book.create_sheet('Histórico');s.append(['Ação','Detalhe','Data UTC'])
    for r in audit:s.append([r['action'],r['detail'],r['created_at']])
    # User strings must remain text, never become executable Excel formulas.
    for s in book:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for cell in s[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='163A60')
        for col in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(col)].width=26
        for row in s:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
    output=BytesIO();book.save(output);return output.getvalue()

def pdf_report(month,rows,prices,contract,closed,goals=None):
    goals=goals or {}
    output=BytesIO();doc=SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=30,leftMargin=30)
    styles=getSampleStyleSheet();story=[]
    def paragraph(text):return Paragraph(escape(str(text)),styles['BodyText'])
    def table(data):
        t=Table(data,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#163a60')),('TEXTCOLOR',(0,0),(-1,0),colors.white),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),9),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.whitesmoke,colors.white]),
            ('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('BOTTOMPADDING',(0,0),(-1,-1),7)]));return t
    story += [Paragraph('Fechamento mensal de leitura de hidrômetros',styles['Title']),paragraph(f"{contract.get('name') or 'Contrato não informado'} | Referência {month} | {'Fechado' if closed else 'PRÉVIA'}"),Spacer(1,14)]
    summary=summarize(rows,prices)
    story.append(table([['Código','Serviço','Quantidade','Unitário','Total']]+[[r['Código'],r['Categoria'],r['Quantidade'],money(r['unit_cents']),money(r['total_cents'])] for r in summary]))
    story += [Spacer(1,14),Paragraph('Total: '+money(sum(r['total_cents'] for r in summary)),styles['Heading1']),paragraph(f"Registros importados: {len(rows)}. Excluídos: {sum(bool(r['excluded']) for r in rows)}. As unidades dos adicionais não representam leituras adicionais."),Spacer(1,12)]
    chart=VerticalBarChart();chart.x=40;chart.y=40;chart.width=620;chart.height=125
    chart.data=[[r['Quantidade'] for r in summary]];chart.categoryAxis.categoryNames=['R$ 1,38','R$ 1,28','R$ 1,37','Impressão','Apontamento'];chart.valueAxis.valueMin=0;chart.bars[0].fillColor=colors.HexColor('#147D92')
    d=Drawing(740,195);d.add(chart);story.append(d)
    for by in ('employee','city'):
        title='Produtividade por funcionário' if by=='employee' else 'Produtividade por unidade / cidade (ATC)'
        story.extend([PageBreak(),Paragraph(title,styles['Heading1'])])
        entries=productivity(rows,prices,by)
        story.append(table([['Funcionário' if by=='employee' else 'ATC','Registros','Leituras','Comunicados','Apontamentos','Valor']]+[[paragraph(next(iter(r.values()))),r['Registros'],r['Leituras faturáveis'],r['Comunicados'],r['Apontamentos'],money(r['total_cents'])] for r in entries]))
    story.extend([PageBreak(),Paragraph('Metas mensais informadas',styles['Heading1'])])
    goal_rows=[]
    for by in ('employee','city'):
        for r in productivity(rows,prices,by):
            name=next(iter(r.values()));goal=goals.get(by+'|'+name)
            if goal:goal_rows.append(['Funcionário' if by=='employee' else 'ATC',paragraph(name),r['Leituras faturáveis'],goal,f"{100*r['Leituras faturáveis']/goal:.1f}%"])
    story.append(table([['Dimensão','Funcionário / ATC','Leituras','Meta','Atingimento']]+goal_rows) if goal_rows else paragraph('Nenhuma meta mensal informada. Metas não são presumidas.'))
    story += [Spacer(1,12),paragraph('Regras: fatura retida não gera serviço de leitura. Causa preenchida gera a categoria R$ 1,37; causa vazia é separada por fatura impressa. Adicionais em todos os registros. CO1 Atual preenchido gera um apontamento, inclusive códigos de vistoria, sem cobrança do serviço de vistoria. Registros sem data e releituras incluídos. O Excel contém a conferência detalhada e os dados importados; originais preservados separadamente.')]
    doc.build(story);return output.getvalue()
