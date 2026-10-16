"""Teste de integração exclusivamente com a planilha de um grupo fornecida."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import openpyxl
import core
import storage
from exports import excel_report, pdf_report, productivity

EXPECTED={'138':15942,'128':3379,'137':142,'impressao':5346,'apontamento':2523}

def validate(content):
    parsed=core.parse_excel(content)
    assert len(parsed)==1
    p=parsed[0]
    assert len(p['headers'])==72 and len(p['records'])==19668
    prices={k:v[2] for k,v in core.CATEGORIES.items()}
    evaluated=[{'raw':r['raw'],'result':core.evaluate(r['raw'],p['mapping'],'2026-10'),'excluded':False} for r in p['records']]
    assert not any(r['result']['errors'] for r in evaluated)
    assert all(sum(r['result']['units'][k] for k in ('138','128','137'))<=1 for r in evaluated)
    assert {r['key']:r['Quantidade'] for r in core.summarize(evaluated,prices)}==EXPECTED
    assert sum(r['total_cents'] for r in core.summarize(evaluated,prices))==3045412
    with tempfile.TemporaryDirectory() as td:
        original_data=storage.DATA;storage.DATA=Path(td)
        try:
            assert storage.import_sheet(content,'planilha.xlsx',p,p['mapping'],'2026-10')==(19668,0)
            try:storage.import_sheet(content,'renomeada.xlsx',p,p['mapping'],'2026-10')
            except ValueError:pass
            else:raise AssertionError('Reimportação não bloqueada')
            archive=next((storage.DATA/'originais').glob('*.xlsx'))
            assert archive.read_bytes()==content
            rows=storage.records('2026-10');assert len(rows)==19668 and not storage.conflicts(rows)
            for dim in ('employee','city'):
                entries=productivity(rows,prices,dim)
                assert sum(r['total_cents'] for r in entries)==3045412
                assert sum(r['Registros'] for r in entries)==19668
            row_id=rows[0]['id'];storage.exclude('2026-10',row_id,'Conferência do registro real')
            assert sum(r['excluded'] for r in storage.records('2026-10'))==1
            storage.exclude('2026-10',row_id,'Reinclusão após conferência',False)
            storage.close_month('2026-10');assert storage.month_info('2026-10')['closed']
            try:storage.update_month_prices('2026-10',{k:0 for k in prices})
            except ValueError:pass
            else:raise AssertionError('Preço do mês fechado alterado')
            storage.close_month('2026-10',False,'Teste de reabertura com dados reais')
            xlsx=excel_report('2026-10',rows,prices,storage.imports('2026-10'),storage.audit('2026-10'),{},False)
            report=openpyxl.load_workbook(BytesIO(xlsx),read_only=True,data_only=True)
            assert report['Dados importados'].max_row==19669 and report['Dados importados'].max_column==73
            assert round(report['Fechamento']['E11'].value*100)==3045412
            report.close()
            pdf=pdf_report('2026-10',rows,prices,{},False);assert pdf.startswith(b'%PDF-') and pdf.rstrip().endswith(b'%%EOF')
        finally:storage.DATA=original_data
    print('PASSOU: 19.668 registros, 72 colunas, cinco categorias, exclusividade, centavos, produtividade, integridade do original, deduplicação, decisão com histórico, fechamento/reabertura, Excel e PDF.')
    print(json.dumps(EXPECTED,ensure_ascii=False));print('TOTAL: R$ 30.454,12')
    return parsed

if __name__=='__main__':
    if len(sys.argv)!=2:raise SystemExit('Uso: python validar_planilha.py caminho_da_planilha.xlsx')
    validate(Path(sys.argv[1]).read_bytes())
