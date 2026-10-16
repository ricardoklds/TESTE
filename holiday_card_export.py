"""Exportação sob demanda dos mesmos IDs cobrados no card; nenhuma gravação na base."""
from collections import Counter
from contextlib import closing
from io import BytesIO
import json
import openpyxl
import holidays
import managerial
from fast_reports import text_row


def prepare(path,revision,month,start,end,filters,expected_quantity,expected_cents):
    book=openpyxl.Workbook(write_only=True)
    try:
        return _prepare(book,path,revision,month,start,end,filters,expected_quantity,expected_cents)
    finally:
        # openpyxl 3.1.5 usa XML temporário no modo incremental. Limpar também ao cancelar.
        # Somente writers pertencentes a este workbook; nunca arquivos do aplicativo.
        from pathlib import Path
        for sheet in book.worksheets:
            try:
                if not sheet.closed:sheet.close()
            finally:
                writer=getattr(sheet,'_writer',None)
                if writer and Path(writer.out).exists():writer.cleanup()


def _prepare(book,path,revision,month,start,end,filters,expected_quantity,expected_cents):
    if expected_quantity<=0:raise ValueError('Não há leituras elegíveis ao adicional de feriado nesta seleção.')
    if expected_cents!=expected_quantity*holidays.UNIT:raise ValueError('Quantidade e valor do card não são consistentes. Atualize o painel.')
    summary=book.create_sheet('Resumo');groups=Counter();count=0;sheet=None
    headers=['Competência','ID do registro','ID da importação','Arquivo original','Aba original','Linha original',
        'Data da leitura original','Data da elegibilidade','Cidade/ATC original','Município na apuração','UF na apuração',
        'Grupo original','Funcionário original','Fornecimento original','Execução original','Código do serviço',
        'Quantidade','Valor unitário (R$)','Valor total (R$)']
    with closing(managerial.connect(path)) as c:
        if c.execute('SELECT revision FROM cache_revision WHERE id=1').fetchone()[0]!=revision:
            raise ValueError('Os dados mudaram. Atualize o painel e solicite a exportação novamente.')
        fee_sql,fee_args=holidays.charges_query(c,month,start,end)
        where,args=managerial.conditions('h',month,start,end,filters)
        where=where.replace('h.month=?','1',1);args=args[1:]
        # Nenhum JOIN de registros contra o calendário; busca de originais por PK em lotes.
        cursor=c.execute('SELECT h.id,h.day,h.municipality,h.state FROM ('+fee_sql+') h WHERE '+where+' ORDER BY h.id',[*fee_args,*args])
        while True:
            batch=cursor.fetchmany(500)
            if not batch:break
            ids=[r['id'] for r in batch]
            originals={r['id']:r for r in c.execute('SELECT r.id,r.import_id,r.line,r.raw,i.name,i.sheet,i.mapping FROM records r LEFT JOIN imports i ON i.id=r.import_id WHERE r.id IN ('+','.join('?' for _ in ids)+')',ids)}
            for row in batch:
                source=originals.get(row['id'])
                if source is None:raise ValueError('Um ID cobrado não possui registro original no banco. Confira a integridade antes de exportar.')
                raw=json.loads(source['raw']);mapping=json.loads(source['mapping']) if source['mapping'] else {}
                get=lambda key:raw.get(mapping.get(key))
                if count%1048575==0:
                    sheet=book.create_sheet('Leituras '+str(count//1048575+1));text_row(sheet,headers)
                text_row(sheet,[month,row['id'],source['import_id'],source['name'],source['sheet'],source['line'],
                    get('date'),row['day'],get('city'),row['municipality'] or None,row['state'] or None,
                    get('group'),get('employee'),get('supply'),get('task'),'200014097',1,holidays.UNIT/100,holidays.UNIT/100])
                count+=1;groups[(get('city'),get('group'),row['municipality'],row['state'])]+=1
        # Mesma transação de leitura: não mistura revisões nem corrige totais para coincidir.
        if count!=expected_quantity:raise ValueError('A consulta detalhada não coincide com o card. Nenhum arquivo será disponibilizado; atualize e confira os filtros.')
        text_row(summary,['Campo','Valor'])
        for pair in [('Competência',month),('Data inicial',start or 'Mês completo'),('Data final',end or 'Mês completo'),
            ('Quantidade',count),('Valor unitário (R$)',holidays.UNIT/100),('Valor total (R$)',count*holidays.UNIT/100),
            ('Revisão',revision),('Filtros',json.dumps(filters,ensure_ascii=False)),('Escopo','Somente IDs efetivamente cobrados; nenhuma nova seleção de duplicidades')]:text_row(summary,list(pair))
        grouped=book.create_sheet('Por cidade e grupo');text_row(grouped,['Cidade/ATC original','Grupo original','Município na apuração','UF','Quantidade','Total (R$)'])
        for key,n in sorted(groups.items(),key=lambda pair:str(pair[0])):text_row(grouped,[*key,n,n*holidays.UNIT/100])
        out=BytesIO();book.save(out)
    return out.getvalue()
