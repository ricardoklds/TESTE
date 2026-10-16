"""Finanças de vistorias em banco separado. Sem dependência das regras de leituras."""
from contextlib import closing
from datetime import datetime, date, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import calendar, hashlib, json, sqlite3, unicodedata, uuid
import inspections
import inspection_contract as contract_policy

SERVICES=(('200014105','Vistoria para aprovação de ligação',2250),('200014106','Vistoria predial',2250),('200014104','Vistoria de atualização cadastral',1200),('geo','Geolocalização',1200))

def norm(value):return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(value or '')) if not unicodedata.combining(c)).upper().split())
def month_bounds(month):
    first=date.fromisoformat(month+'-01');return first.isoformat(),first.replace(day=calendar.monthrange(first.year,first.month)[1]).isoformat()
def now():return datetime.now(timezone.utc).isoformat()
def digest(value):return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def backup(path):
    path=Path(path)
    if not path.exists():return None
    dest=path.parent/'backups';dest.mkdir(exist_ok=True);target=dest/(path.stem+'-antes-financeiro-'+uuid.uuid4().hex+'.sqlite3')
    with closing(sqlite3.connect(path.as_uri()+'?mode=ro',uri=True)) as src,closing(sqlite3.connect(target)) as dst:
        src.backup(dst)
        if dst.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise ValueError('Backup inválido: operação cancelada.')
    return str(target)

def connect(root):
    root=Path(root).resolve();root.mkdir(parents=True,exist_ok=True);p=root/'financeiro.sqlite3'
    c=sqlite3.connect(p.as_uri(),uri=True,factory=inspections.ClosingConnection);c.row_factory=sqlite3.Row;c.execute('PRAGMA busy_timeout=30000')
    version=c.execute('PRAGMA user_version').fetchone()[0]
    if version not in (0,1,2):c.close();raise ValueError('Banco financeiro desconhecido; migração não autorizada.')
    if version in (1,2) and c.execute("SELECT 1 FROM metadata WHERE key='catalog_initial'").fetchone():
        if version==1:
            try:_migrate_contract(c,root)
            except Exception:c.close();raise
        return c
    c.executescript('''CREATE TABLE IF NOT EXISTS revision(id INTEGER PRIMARY KEY CHECK(id=1),value INTEGER NOT NULL); INSERT OR IGNORE INTO revision VALUES(1,0);
    CREATE TABLE IF NOT EXISTS services(code TEXT PRIMARY KEY,name TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS prices(code TEXT NOT NULL,effective TEXT NOT NULL,cents INTEGER NOT NULL CHECK(cents>=0),active INTEGER NOT NULL,PRIMARY KEY(code,effective));
    CREATE TABLE IF NOT EXISTS rules(id INTEGER PRIMARY KEY,tss TEXT NOT NULL,cause TEXT NOT NULL DEFAULT '',form TEXT NOT NULL DEFAULT '',effective TEXT NOT NULL,code TEXT,decision TEXT NOT NULL,evidence TEXT NOT NULL,UNIQUE(tss,cause,form,effective));
    CREATE TABLE IF NOT EXISTS commercial(id INTEGER PRIMARY KEY,day TEXT,city TEXT NOT NULL,os TEXT NOT NULL,contract TEXT NOT NULL,tss TEXT NOT NULL,cause TEXT NOT NULL,form TEXT NOT NULL,family TEXT NOT NULL,closed INTEGER NOT NULL,inactive INTEGER NOT NULL,raw_hash TEXT NOT NULL);
    CREATE INDEX IF NOT EXISTS commercial_period ON commercial(day,closed,tss);
    CREATE INDEX IF NOT EXISTS commercial_os ON commercial(contract,os);
    CREATE TABLE IF NOT EXISTS reviews(visit_id INTEGER NOT NULL,month TEXT NOT NULL,decision TEXT NOT NULL,code TEXT,evidence TEXT NOT NULL,PRIMARY KEY(visit_id,month));
    CREATE TABLE IF NOT EXISTS geo(month TEXT PRIMARY KEY,mode TEXT NOT NULL,total INTEGER,cities TEXT NOT NULL,evidence TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS approvals(month TEXT PRIMARY KEY,payload TEXT NOT NULL,created TEXT NOT NULL,evidence TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,created TEXT NOT NULL,action TEXT NOT NULL,payload TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
    PRAGMA user_version=1;''')
    for code,name,cents in SERVICES:
        c.execute('INSERT OR IGNORE INTO services VALUES(?,?)',(code,name));c.execute('INSERT OR IGNORE INTO prices VALUES(?,?,?,1)',(code,'2026-09',cents))
    if not c.execute("SELECT 1 FROM metadata WHERE key='catalog_initial'").fetchone():
        from inspection_catalog import MAPPING,EXCLUDED
        evidence='Mapeamento inicial confirmado pelo usuário em 15/10/2026; sem inferência por cores ou similaridade.'
        for code,descriptions in MAPPING.items():
            for description in descriptions:c.execute('INSERT OR IGNORE INTO rules(tss,effective,code,decision,evidence) VALUES(?,?,?,?,?)',(norm(description),'2026-09',code,'Cobrar',evidence))
        for description in EXCLUDED:c.execute('INSERT OR IGNORE INTO rules(tss,effective,decision,evidence) VALUES(?,?,?,?)',(norm(description),'2026-09','Excluir',evidence))
        c.execute("INSERT INTO metadata VALUES('catalog_initial','1')")
        audit(c,'Catálogo inicial confirmado',{'evidence':evidence})
    c.commit();_migrate_contract(c,root);return c

def revision(root):
    with connect(root) as c:return c.execute('SELECT value FROM revision WHERE id=1').fetchone()[0]
def audit(c,action,payload):
    c.execute('INSERT INTO audit(created,action,payload) VALUES(?,?,?)',(now(),action,json.dumps(payload,ensure_ascii=False)))
    c.execute('UPDATE revision SET value=value+1 WHERE id=1')
def editable(c,month):
    month_bounds(month)
    if c.execute('SELECT 1 FROM approvals WHERE month=?',(month,)).fetchone():raise ValueError('Competência aprovada; histórico preservado.')
def future(c,effective):
    month_bounds(effective)
    if c.execute('SELECT 1 FROM approvals WHERE month>=? LIMIT 1',(effective,)).fetchone():raise ValueError('Vigência alcança competência já aprovada. Use mês posterior.')

def sync(root,progress=None):
    """Projeção incremental, explícita e transacional. Não modifica banco operacional."""
    root=Path(root).resolve();source=root/'vistorias.sqlite3'
    if not source.exists():return 0
    # Backup consistente antes de preparar o índice financeiro da base existente.
    with connect(root) as check:prepared=check.execute("SELECT 1 FROM metadata WHERE key='prepared'").fetchone()
    if not prepared:backup(source);backup(root/'financeiro.sqlite3')
    with closing(sqlite3.connect(source.as_uri()+'?mode=ro',uri=True)) as src,connect(root) as c:
        c.execute('BEGIN IMMEDIATE');last=c.execute('SELECT COALESCE(MAX(id),0) FROM commercial').fetchone()[0];n=0
        cursor=src.execute('SELECT id,day,city,os,raw FROM visits WHERE id>? ORDER BY id',(last,))
        while batch:=cursor.fetchmany(2000):
            values=[]
            for id,day,city,os,rawtext in batch:
                r=json.loads(rawtext);closed=all(norm(r.get(k))=='FECHADA' for k in ('Status da OS','Status da Operação'))
                values.append((id,day,city,os,norm(r.get('Contrato')),norm(r.get('TSS')),norm(r.get('Causa Resultado')),norm(r.get('Formulário de Baixa')),norm(r.get('Família')),int(closed),int(norm(r.get('TSS'))=='VISTORIA LIGACAO INATIVA SEM DEBITO'),hashlib.sha256(rawtext.encode()).hexdigest(),norm(city),execution_stamp(r.get('Data Fim Execução'),day),day[:7] if day else None))
            c.executemany('INSERT INTO commercial VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',values);n+=len(values)
            if progress:progress(n)
        c.execute("INSERT OR REPLACE INTO metadata VALUES('prepared','1')")
        if n:audit(c,'Atualizar projeção comercial',{'rows':n,'last_id':last})
    return n

def service_list(root,month):
    with connect(root) as c:return [dict(r) for r in c.execute('''SELECT s.code,s.name,p.effective,p.cents,p.active FROM services s JOIN prices p ON p.code=s.code
      WHERE p.effective=(SELECT MAX(x.effective) FROM prices x WHERE x.code=s.code AND x.effective<=?) ORDER BY s.code''',(month,))]

def save_service(root,code,name,effective,cents,active,evidence):
    code=code.strip()
    if not code or not name.strip() or type(cents)!=int or cents<0 or not evidence.strip():raise ValueError('Código contratual confirmado, descrição, preço e justificativa obrigatórios.')
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE');future(c,effective)
        existing=c.execute('SELECT name FROM services WHERE code=?',(code,)).fetchone()
        if existing and existing[0]!=name.strip():raise ValueError('Preserve a descrição do código existente; adicione outro serviço quando necessário.')
        before=c.execute('SELECT * FROM prices WHERE code=? AND effective=?',(code,effective)).fetchone()
        c.execute('INSERT OR IGNORE INTO services VALUES(?,?)',(code.strip(),name.strip()))
        c.execute('INSERT OR REPLACE INTO prices VALUES(?,?,?,?)',(code.strip(),effective,cents,int(active)))
        audit(c,'Preço/vigência',{'code':code,'name':name,'effective':effective,'cents':cents,'active':bool(active),'evidence':evidence,'before':dict(before) if before else None})

def save_rule(root,tss,cause,form,effective,decision,code,evidence):
    if decision not in ('Cobrar','Excluir','Pendente') or not norm(tss) or not evidence.strip():raise ValueError('TSS, decisão e comprovação obrigatórios.')
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE');future(c,effective)
        if decision=='Cobrar' and (code=='geo' or not c.execute('SELECT 1 FROM services WHERE code=?',(code,)).fetchone()):raise ValueError('Serviço contratual de vistoria inválido.')
        payload={'tss':norm(tss),'cause':norm(cause),'form':norm(form),'effective':effective,'decision':decision,'code':code if decision=='Cobrar' else None,'evidence':evidence}
        before=c.execute('SELECT * FROM rules WHERE tss=? AND cause=? AND form=? AND effective=?',(payload['tss'],payload['cause'],payload['form'],effective)).fetchone()
        c.execute('INSERT OR REPLACE INTO rules(tss,cause,form,effective,code,decision,evidence) VALUES(:tss,:cause,:form,:effective,:code,:decision,:evidence)',payload)
        audit(c,'Classificação TSS',{**payload,'before':dict(before) if before else None})

def review(root,visit_id,month,decision,code,evidence):
    if decision not in ('Cobrar','Excluir') or not evidence.strip():raise ValueError('Decisão e justificativa obrigatórias.')
    first,last=month_bounds(month)
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE');editable(c,month)
        if not c.execute('SELECT 1 FROM commercial WHERE id=? AND day BETWEEN ? AND ?',(visit_id,first,last)).fetchone():raise ValueError('Registro não pertence ao mês informado.')
        if decision=='Cobrar' and (code=='geo' or not c.execute('SELECT 1 FROM services WHERE code=?',(code,)).fetchone()):raise ValueError('Serviço inválido.')
        before=c.execute('SELECT * FROM reviews WHERE visit_id=? AND month=?',(visit_id,month)).fetchone()
        c.execute('INSERT OR REPLACE INTO reviews VALUES(?,?,?,?,?)',(visit_id,month,decision,code if decision=='Cobrar' else None,evidence))
        audit(c,'Conferência individual',{'id':visit_id,'month':month,'decision':decision,'code':code,'evidence':evidence,'before':dict(before) if before else None})

def save_geo(root,month,mode,total,cities,evidence=None):
    if mode not in ('Total mensal','Por município'):raise ValueError('Modo de lançamento inválido.')
    evidence=(evidence or '').strip() or 'Lançamento manual pela interface; documento não exigido.'
    if total is not None and (type(total)!=int or total<0):raise ValueError('Quantidade total inválida.')
    if any(not k.strip() or type(v)!=int or v<0 for k,v in cities.items()):raise ValueError('Município e quantidades inteiras não negativas obrigatórios.')
    names=[norm(k) for k in cities]
    if len(names)!=len(set(names)):raise ValueError('Município repetido.')
    if (mode=='Total mensal' and total is None) or (mode=='Por município' and not cities):raise ValueError('Preencha o modo escolhido; informe zero explicitamente quando não houver serviço.')
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE');editable(c,month)
        before=c.execute('SELECT * FROM geo WHERE month=?',(month,)).fetchone()
        c.execute('INSERT OR REPLACE INTO geo VALUES(?,?,?,?,?)',(month,mode,total,json.dumps(cities,ensure_ascii=False),evidence))
        audit(c,'Geolocalização',{'month':month,'before':dict(before) if before else None,'mode':mode,'total':total,'cities':cities,'evidence':evidence})

def geo_summary(c,month,prices):
    r=c.execute('SELECT * FROM geo WHERE month=?',(month,)).fetchone()
    if not r:return {'quantity':0,'unit_cents':prices['geo']['cents'] if 'geo' in prices else None,'cents':0,'ready':False,'reason':'Geolocalização não confirmada (zero deve ser informado explicitamente).'}
    return geo_values(r,contract_policy.policy(c,month),prices.get('geo'))

def geo_values(r,policy,price):
    """Prévia em memória; o mesmo critério do lançamento salvo, sem gravar."""
    cities=r['cities'] if isinstance(r['cities'],dict) else json.loads(r['cities'])
    inside={k:v for k,v in cities.items() if contract_policy.classify(k,policy)=='Dentro do escopo'}
    unknown={k:v for k,v in cities.items() if contract_policy.classify(k,policy).startswith('Pendente') and v}
    outside=sum(v for k,v in cities.items() if contract_policy.classify(k,policy)=='Fora do escopo do contrato')
    quantity=r['total'] if r['mode']=='Total mensal' else sum(inside.values())
    matches=r['total'] is None or not cities or r['total']==sum(inside.values())
    ready=matches and not unknown and price is not None and bool(price['active'])
    return {'quantity':quantity,'unit_cents':price['cents'] if price else None,'cents':quantity*price['cents'] if price else 0,'ready':ready,'mode':r['mode'],'cities_total':sum(inside.values()) if cities else None,'reference_total':r['total'],'out_of_scope_quantity':outside,'unknown_municipalities':unknown,'reason':'' if ready else 'Total e municípios contratuais divergem, município pendente ou preço inativo/ausente.'}

# Todas as classificações em SQL; registros brutos são lidos só na projeção/importação e exportação solicitada.
def classified_sql(month,connection=None):
    month_bounds(month)
    periods={month}
    if connection is not None:periods.update(r[0] for r in connection.execute('SELECT DISTINCT month FROM commercial WHERE month IS NOT NULL'))
    contexts=[{'month':m,**contract_policy.policy(connection,m)} for m in sorted(periods)]
    sql='''WITH policies AS (SELECT json_extract(value,'$.month') month,json_extract(value,'$.allowed') allowed,json_extract(value,'$.excluded') excluded,json_extract(value,'$.modes') modes,json_extract(value,'$.legacy') legacy FROM json_each(:contexts)),
    period_keys AS MATERIALIZED (SELECT DISTINCT contract,os FROM commercial WHERE month=:month),
    source AS MATERIALIZED (SELECT * FROM commercial WHERE month=:month UNION ALL SELECT v.* FROM period_keys k JOIN commercial v ON v.contract=k.contract AND v.os=k.os WHERE v.month<>:month),
    allowed AS MATERIALIZED (SELECT p.month,j.value city FROM policies p,json_each(p.allowed) j),
    excluded AS MATERIALIZED (SELECT p.month,j.value city FROM policies p,json_each(p.excluded) j),
    effective_rules AS MATERIALIZED (SELECT pol.month,r.* FROM policies pol JOIN rules r ON r.effective=(SELECT MAX(effective) FROM rules x WHERE x.tss=r.tss AND x.cause=r.cause AND x.form=r.form AND x.effective<=pol.month)),
    effective_prices AS MATERIALIZED (SELECT pol.month,p.* FROM policies pol JOIN prices p ON p.effective=(SELECT MAX(effective) FROM prices x WHERE x.code=p.code AND x.effective<=pol.month)),
    matches AS (SELECT v.id,r.code,r.decision,ROW_NUMBER() OVER(PARTITION BY v.id ORDER BY (r.cause<>'')+(r.form<>'') DESC,r.id DESC) priority,
      COUNT(*) OVER(PARTITION BY v.id,(r.cause<>'')+(r.form<>'')) ambiguity FROM source v JOIN effective_rules r ON r.month=v.month AND r.tss=v.tss AND (r.cause='' OR r.cause=v.cause) AND (r.form='' OR r.form=v.form)),
    override_values AS (SELECT json_extract(value,'$.tss') tss,json_extract(value,'$.effective') effective,json_extract(value,'$.cents') cents FROM metadata WHERE key LIKE 'tss_price:%'),
    overrides AS (SELECT id,cents FROM (SELECT v.id,o.cents,ROW_NUMBER() OVER(PARTITION BY v.id ORDER BY o.effective DESC) precedence FROM source v JOIN override_values o ON o.tss=v.tss AND o.effective<=v.month) WHERE precedence=1),
    chosen AS MATERIALIZED (SELECT v.*,COALESCE(rv.decision,m.decision) decision,COALESCE(rv.code,m.code) code,rv.visit_id reviewed,COALESCE(m.ambiguity,0) ambiguity,
      CASE WHEN pol.legacy THEN 'Dentro do escopo' WHEN a.city IS NOT NULL THEN 'Dentro do escopo' WHEN e.city IS NOT NULL THEN 'Fora do escopo do contrato' ELSE 'Pendente: município não autorizado' END scope,
      pol.legacy,COALESCE(json_extract(pol.modes,'$.'||'"'||COALESCE(rv.code,m.code)||'"'),'manual') mode
      FROM source v JOIN policies pol ON pol.month=v.month LEFT JOIN allowed a ON a.month=v.month AND a.city=v.city_key LEFT JOIN excluded e ON e.month=v.month AND e.city=v.city_key LEFT JOIN matches m ON m.id=v.id AND m.priority=1 LEFT JOIN reviews rv ON rv.visit_id=v.id AND rv.month=v.month),
    manual_counts AS MATERIALIZED (SELECT month,contract,code,os,SUM(scope='Dentro do escopo' AND closed=1) n FROM chosen WHERE mode='manual' GROUP BY month,contract,code,os),
    legacy_counts AS MATERIALIZED (SELECT month,contract,os,COUNT(*) n FROM chosen WHERE legacy=1 GROUP BY month,contract,os),
    execution_counts AS MATERIALIZED (SELECT contract,code,os,executed_at,city_key,COUNT(*) n FROM chosen WHERE mode='execution' AND scope='Dentro do escopo' AND closed=1 AND decision='Cobrar' GROUP BY contract,code,os,executed_at,city_key),
    base AS MATERIALIZED (SELECT v.*,COALESCE(o.cents,p.cents) cents,p.active,COALESCE(mc.n,0) repeated,COALESCE(lc.n,0) legacy_repeated,COALESCE(ec.n,0) execution_repeated,
      CASE WHEN scope='Fora do escopo do contrato' THEN scope WHEN closed=0 THEN 'Não concluída' WHEN scope<>'Dentro do escopo' THEN scope
      WHEN decision='Excluir' AND (reviewed IS NOT NULL OR ambiguity<=1) THEN 'Excluído por classificação confirmada'
      WHEN reviewed IS NULL AND inactive THEN 'Pendente: ligação inativa'
      WHEN reviewed IS NULL AND ambiguity>1 THEN 'Pendente: regras ambíguas'
      WHEN decision='Pendente' THEN 'Pendente: classificação definida pelo usuário'
      WHEN v.executed_at IS NULL THEN 'Pendente: data de execução sem horário válido' WHEN v.os='' THEN 'Pendente: OS ausente' WHEN decision IS NULL THEN 'Pendente: serviço não mapeado'
      WHEN p.cents IS NULL OR p.active<>1 THEN 'Pendente: preço ausente ou serviço inativo' WHEN mode='once' AND EXISTS(SELECT 1 FROM frozen_charges f WHERE f.contract=v.contract AND f.code=v.code AND f.os=v.os AND f.month<>v.month) THEN 'Pendente: OS já faturada em competência aprovada' ELSE 'Elegível' END base_reason
      FROM chosen v LEFT JOIN effective_prices p ON p.month=v.month AND p.code=v.code LEFT JOIN overrides o ON o.id=v.id
      LEFT JOIN manual_counts mc ON mc.month=v.month AND mc.contract=v.contract AND mc.code=v.code AND mc.os=v.os
      LEFT JOIN legacy_counts lc ON lc.month=v.month AND lc.contract=v.contract AND lc.os=v.os
      LEFT JOIN execution_counts ec ON ec.contract=v.contract AND ec.code=v.code AND ec.os=v.os AND ec.executed_at=v.executed_at AND ec.city_key=v.city_key),
    ranked AS (SELECT id,contract,code,os,executed_at,city,month,ROW_NUMBER() OVER(PARTITION BY contract,code,os ORDER BY executed_at DESC,id ASC) winner_rank,
      COUNT(*) OVER(PARTITION BY contract,code,os) os_candidates,
      MIN(city_key) OVER(PARTITION BY contract,code,os) city_min,MAX(city_key) OVER(PARTITION BY contract,code,os) city_max,
      FIRST_VALUE(id) OVER(PARTITION BY contract,code,os ORDER BY executed_at DESC,id ASC) winner_id
      FROM base WHERE base_reason='Elegível' AND mode='once'),
    final AS (SELECT v.*,r.winner_id,r.os_candidates,r.executed_at winner_executed_at,r.city winner_city,r.month winner_month,CASE WHEN v.base_reason<>'Elegível' THEN v.base_reason
      WHEN v.mode='once' AND EXISTS(SELECT 1 FROM frozen_charges f WHERE f.contract=v.contract AND f.code=v.code AND f.os=v.os AND f.month<>v.month) THEN 'Pendente: OS já faturada em competência aprovada'
      WHEN v.mode='once' AND r.city_min<>r.city_max THEN 'Pendente: OS com municípios divergentes'
      WHEN v.mode='once' AND r.winner_rank>1 THEN 'Sem cobrança adicional — OS já representada'
      WHEN v.mode='execution' AND v.execution_repeated>1 AND v.reviewed IS NULL THEN 'Pendente: execução sem identidade unívoca'
      WHEN v.mode='manual' AND (CASE WHEN v.legacy THEN v.legacy_repeated ELSE v.repeated END)>1 AND v.reviewed IS NULL THEN 'Pendente: OS repetida'
      ELSE 'Faturável' END reason FROM base v LEFT JOIN ranked r ON r.id=v.id)
    SELECT * FROM final WHERE month=:month'''
    if len(periods)==1:
        start=sql.index('    period_keys AS');end=sql.index('    allowed AS',start)
        sql=sql[:start]+'    source AS MATERIALIZED (SELECT * FROM commercial WHERE month=:month),\n'+sql[end:]
    return sql,{'month':month,'contexts':json.dumps(contexts,ensure_ascii=False)}


def preview(root,month,connection=None):
    if connection is None:
        with connect(root) as c:return preview(root,month,c)
    c=connection;sql,args=classified_sql(month,c)
    municipal_groups=[dict(r) for r in c.execute("SELECT reason,code,cents unit_cents,city_key,scope,COUNT(*) quantity,SUM(closed=1) closed_quantity,COALESCE(SUM(CASE WHEN reason='Faturável' THEN cents ELSE 0 END),0) cents FROM ("+sql+") GROUP BY reason,code,cents,city_key,scope",args)]
    units={};municipalities={}
    for r in municipal_groups:
        key=(r['reason'],r['code'],r['unit_cents']);u=units.setdefault(key,{'reason':r['reason'],'code':r['code'],'unit_cents':r['unit_cents'],'quantity':0,'cents':0});u['quantity']+=r['quantity'];u['cents']+=r['cents']
        if r['scope']=='Fora do escopo do contrato':continue
        m=municipalities.setdefault(r['city_key'],{'municipality':r['city_key'],'operational_count':0,'billed_count':0,'cents':0,'pending':0})
        if r['scope']=='Dentro do escopo':m['operational_count']+=r['closed_quantity']
        if r['reason']=='Faturável':m['billed_count']+=r['quantity'];m['cents']+=r['cents']
        if r['reason'].startswith('Pendente:'):m['pending']+=r['quantity']
    unit_groups=list(units.values())
    combined={}
    for r in unit_groups:
        key=(r['reason'],r['code']);g=combined.setdefault(key,{'reason':r['reason'],'code':r['code'],'quantity':0,'cents':0});g['quantity']+=r['quantity'];g['cents']+=r['cents']
    groups=list(combined.values())
    # Todas as linhas sem data ficam pendentes globais: não atribuir mês por suposição.
    pol=contract_policy.policy(c,month)
    undated=sum(r[1] for r in c.execute('SELECT city_key,COUNT(*) FROM commercial WHERE day IS NULL AND closed=1 GROUP BY city_key') if contract_policy.classify(r[0],pol)!='Fora do escopo do contrato')
    pending=sum(r['quantity'] for r in groups if r['reason'].startswith('Pendente:'))+undated
    services=service_list_connection(c,month);geo=geo_summary(c,month,services)
    projection_last=c.execute('SELECT COALESCE(MAX(id),0) FROM commercial').fetchone()[0]
    visits=Path(root)/'vistorias.sqlite3';actual_last=0
    if visits.exists():
        with closing(sqlite3.connect(visits.resolve().as_uri()+'?mode=ro',uri=True)) as op:actual_last=op.execute('SELECT COALESCE(MAX(id),0) FROM visits').fetchone()[0]
    ready=projection_last==actual_last and pending==0
    rates=[{k:r[k] for k in ('code','unit_cents','quantity')} for r in unit_groups if r['reason']=='Faturável']
    byservice=[{'code':code,'name':p['name'],'quantity':sum(r['quantity'] for r in groups if r['reason']=='Faturável' and r['code']==code),'cents':sum(r['cents'] for r in groups if r['reason']=='Faturável' and r['code']==code),'unit_cents':next((r['unit_cents'] for r in rates if r['code']==code),p['cents']) if len([r for r in rates if r['code']==code])<=1 else None} for code,p in services.items() if code!='geo']
    result={'month':month,'groups':groups,'municipalities':[municipalities[k] for k in sorted(municipalities)],'policy':pol,'services':byservice,'rates':rates,'pending':pending,'undated':undated,'ready':ready,'sync_pending':actual_last-projection_last,'geo':geo,'inspection_cents':sum(r['cents'] for r in groups),'revision':c.execute('SELECT value FROM revision WHERE id=1').fetchone()[0]}
    result['fingerprint']=digest(result);return result

def service_list_connection(c,month):
    return {r['code']:dict(r) for r in c.execute('''SELECT s.code,s.name,p.* FROM services s JOIN prices p ON p.code=s.code WHERE p.effective=(SELECT MAX(effective) FROM prices x WHERE x.code=s.code AND x.effective<=?)''',(month,))}

def approve(root,month,expected,evidence,confirmed=False):
    if not confirmed or not evidence.strip():raise ValueError('Conciliação e confirmação explícita obrigatórias.')
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE');editable(c,month);p=preview(root,month,c)
        if p['fingerprint']!=expected:raise ValueError('Dados alterados; gere novamente a conciliação.')
        if not p['ready'] or not p['geo']['ready']:raise ValueError('Há pendências de classificação, sincronização ou geolocalização.')
        first,last=month_bounds(month)
        total,municipal,pending=inspections.summary(root,first,last)
        p['deadline_snapshot']={'total':total,'municipalities':municipal,'pending':pending}
        sql,args=classified_sql(month,c)
        c.execute("INSERT INTO frozen_charges(month,visit_id,contract,code,os,cents) SELECT month,id,contract,code,os,cents FROM ("+sql+") WHERE reason='Faturável'",args)
        c.execute('INSERT INTO approvals VALUES(?,?,?,?)',(month,json.dumps(p,ensure_ascii=False),now(),evidence));audit(c,'Aprovar faturamento de vistorias',{'month':month,'snapshot':p,'evidence':evidence})
    return p

def financial(root,month):
    with connect(root) as c:
        r=c.execute('SELECT payload FROM approvals WHERE month=?',(month,)).fetchone()
        if r:return {**json.loads(r[0]),'approved':True}
    return {**preview(root,month),'approved':False}

def traced_sql(c,root,sql):
    """Acrescenta proveniência sem copiar registros brutos para a memória."""
    source=Path(root).resolve()/'vistorias.sqlite3'
    if not source.exists():return 'SELECT v.*,NULL original_file,NULL original_sheet,NULL original_line,NULL original_import,NULL representative_file,NULL representative_sheet,NULL representative_line FROM ('+sql+') v'
    c.execute('ATTACH DATABASE ? AS original',(source.as_uri()+'?mode=ro',))
    return 'SELECT v.*,i.name original_file,i.sheet original_sheet,o.line original_line,o.import_id original_import,wi.name representative_file,wi.sheet representative_sheet,w.line representative_line FROM ('+sql+') v LEFT JOIN original.visits o ON o.id=v.id LEFT JOIN original.imports i ON i.id=o.import_id LEFT JOIN original.visits w ON w.id=v.winner_id LEFT JOIN original.imports wi ON wi.id=w.import_id'

def detail_page(root,month,offset=0,limit=50):
    if not 1<=limit<=100 or offset<0:raise ValueError('Paginação inválida.')
    with connect(root) as c:
        sql,args=classified_sql(month,c);args.update(limit=limit,offset=offset)
        sql='SELECT * FROM ('+sql+') ORDER BY id LIMIT :limit OFFSET :offset'
        return [dict(r) for r in c.execute(traced_sql(c,root,sql),args)]

def reconciliation(root,month,references=None,detail=False):
    """Referências são somente comparação, nunca participam da classificação."""
    from io import BytesIO
    import openpyxl
    from fast_reports import text_row
    p=preview(root,month);w=openpyxl.Workbook(write_only=True);s=w.create_sheet('Conciliação');text_row(s,['Código','Serviço','Calculado dos registros','Referência informada','Diferença','Valor apurado em centavos'])
    for r in p['services']:
        ref=(references or {}).get(r['code']);text_row(s,[r['code'],r['name'],r['quantity'],ref,r['quantity']-ref if ref is not None else None,r['cents']])
    s=w.create_sheet('Motivos');text_row(s,['Motivo','Código','Quantidade','Valor em centavos'])
    for r in p['groups']:text_row(s,[r['reason'],r['code'],r['quantity'],r['cents']])
    text_row(s,['Sem data de execução (globais)',None,p['undated'],0])
    s=w.create_sheet('Metadados');text_row(s,['Competência',month]);text_row(s,['Impressão digital da prévia',p['fingerprint']]);text_row(s,['Situação','Prévia; depende de aprovação explícita']);text_row(s,['Regra','OS e operação Fechada; competência por execução; sem dependência da tabela dinâmica. Referências não afetam cálculos.'])
    with connect(root) as ctx:sql,args=classified_sql(month,ctx)
    s=w.create_sheet('Apuração por TSS');text_row(s,['TSS','Código','Motivo','Quantidade','Valor em centavos'])
    with connect(root) as c:
        for r in c.execute("SELECT tss,code,reason,COUNT(*),SUM(CASE WHEN reason='Faturável' THEN cents ELSE 0 END) FROM ("+sql+") GROUP BY tss,code,reason ORDER BY code,tss",args):text_row(s,list(r))
    if detail:
        with connect(root) as ctx:sql,args=classified_sql(month,ctx)
        with connect(root) as c:
            n=0
            for r in c.execute(traced_sql(c,root,sql)+' ORDER BY v.id',args):
                if n%1048575==0:
                    s=w.create_sheet('Registros '+str(1+n//1048575));text_row(s,['ID','OS','Data','Município','TSS','Causa','Formulário','Código','Motivo','Valor em centavos','Arquivo','Aba','Linha','Importação','ID representante','Arquivo representante','Aba representante','Linha representante','Execução representante','Município representante','Competência representante'])
                text_row(s,[r['id'],r['os'],r['day'],r['city'],r['tss'],r['cause'],r['form'],r['code'],r['reason'],r['cents'] if r['reason']=='Faturável' else 0,r['original_file'],r['original_sheet'],r['original_line'],r['original_import'],r['winner_id'],r['representative_file'],r['representative_sheet'],r['representative_line'],r['winner_executed_at'],r['winner_city'],r['winner_month']]);n+=1
    b=BytesIO();w.save(b);return b.getvalue()

def export_catalog(root,month,include_frozen=True):
    """Configuração vigente para o recálculo independente, sem acessar Excel."""
    with connect(root) as c:
        rules=[dict(r) for r in c.execute('SELECT * FROM rules r WHERE effective=(SELECT MAX(effective) FROM rules x WHERE x.tss=r.tss AND x.cause=r.cause AND x.form=r.form AND effective<=?) AND r.effective<=?',(month,month))]
        reviews={r['raw_hash']:{'decision':r['decision'],'code':r['code']} for r in c.execute('SELECT v.raw_hash,r.decision,r.code FROM reviews r JOIN commercial v ON v.id=r.visit_id WHERE r.month=?',(month,))}
        overrides={}
        for r in c.execute("SELECT value FROM metadata WHERE key LIKE 'tss_price:%'"):
            o=json.loads(r[0])
            if o['effective']<=month and (o['tss'] not in overrides or o['effective']>overrides[o['tss']]['effective']):overrides[o['tss']]=o
        return {'month':month,'policy':contract_policy.policy(c,month),'contexts':{r[0]:contract_policy.policy(c,r[0]) for r in c.execute('SELECT DISTINCT month FROM commercial WHERE month IS NOT NULL')},'frozen_charges':([dict(r) for r in c.execute('SELECT DISTINCT f.month,f.contract,f.code,f.os FROM (SELECT DISTINCT contract,os FROM commercial WHERE month=?) k JOIN frozen_charges f ON f.contract=k.contract AND f.os=k.os',(month,))] if include_frozen else []),'prices':service_list_connection(c,month),'rules':rules,'reviews':reviews,'tss_prices':overrides,'scope':'Catálogo e decisões exportados para '+month}


def catalog_state(root):
    """Catálogo pequeno; não lê registros individuais para preencher o formulário."""
    with connect(root) as c:
        rules=[dict(r) for r in c.execute('SELECT * FROM rules ORDER BY effective,id')]
        custom=[json.loads(r[0]) for r in c.execute("SELECT value FROM metadata WHERE key LIKE 'tss_price:%'")]
        return {'revision':c.execute('SELECT value FROM revision WHERE id=1').fetchone()[0],'rules':rules,'tss_prices':custom}


def _write_change(c,change):
    """Aplicação comum à simulação e à confirmação; sem commits intermediários."""
    kind=change.get('kind');effective=change.get('effective','');future(c,effective)
    evidence=change.get('evidence','').strip()
    if not evidence:raise ValueError('Documento/justificativa obrigatório.')
    if kind=='services':
        children=change.get('changes',[])
        if not children or len({(r.get('code'),r.get('effective')) for r in children})!=len(children):raise ValueError('Lote vazio ou serviços repetidos na mesma vigência.')
        if effective!=min(r.get('effective','') for r in children):raise ValueError('Vigência do lote inconsistente.')
        return {'services':[_write_change(c,{**r,'kind':'service','evidence':evidence}) for r in children]}
    if kind=='service':
        code=change.get('code','').strip();name=change.get('name','').strip();cents=change.get('cents')
        if not code or not name or type(cents)!=int or cents<0:raise ValueError('Código confirmado, descrição e preço válidos obrigatórios.')
        old=c.execute('SELECT name FROM services WHERE code=?',(code,)).fetchone()
        if old and old[0]!=name:raise ValueError('Preserve a descrição do código existente.')
        before=c.execute('SELECT * FROM prices WHERE code=? AND effective=?',(code,effective)).fetchone()
        c.execute('INSERT OR IGNORE INTO services VALUES(?,?)',(code,name))
        c.execute('INSERT OR REPLACE INTO prices VALUES(?,?,?,?)',(code,effective,cents,int(change.get('active',True))))
        return {'price':dict(before) if before else None}
    if kind in ('scope','billing_mode'):
        if kind=='scope':
            allowed=sorted({norm(v) for v in change.get('allowed',[]) if norm(v)})
            excluded=sorted({norm(v) for v in change.get('excluded',[]) if norm(v)})
            if not allowed or set(allowed)&set(excluded):raise ValueError('Informe municípios autorizados, sem sobreposição com os excluídos.')
            key='contract_scope:'+effective;value={'effective':effective,'allowed':allowed,'excluded':excluded}
        else:
            code=change.get('code');mode=change.get('mode')
            if mode not in contract_policy.LABELS or code=='geo' or not c.execute('SELECT 1 FROM services WHERE code=?',(code,)).fetchone():raise ValueError('Categoria/critério inválido.')
            key='billing_mode:'+code+':'+effective;value={'effective':effective,'code':code,'mode':mode,'representative':'latest'}
        old=c.execute('SELECT value FROM metadata WHERE key=?',(key,)).fetchone()
        c.execute('INSERT OR REPLACE INTO metadata VALUES(?,?)',(key,json.dumps(value,ensure_ascii=False)))
        return {'previous_version':json.loads(old[0]) if old else None}
    if kind!='rule':raise ValueError('Tipo de alteração inválido.')
    tss=norm(change.get('tss'));cause=norm(change.get('cause'));form=norm(change.get('form'));decision=change.get('decision');code=change.get('code')
    if not tss or decision not in ('Cobrar','Excluir','Pendente'):raise ValueError('TSS e situação válidos obrigatórios.')
    if decision=='Cobrar' and (code=='geo' or not c.execute('SELECT 1 FROM services WHERE code=?',(code,)).fetchone()):raise ValueError('Categoria contratual inválida.')
    amount=change.get('unit_cents')
    if amount is not None and (type(amount)!=int or amount<0):raise ValueError('Preço específico deve ser inteiro não negativo em centavos.')
    if amount is not None and (cause or form):raise ValueError('Preço por TSS abrange a descrição inteira. Para diferenciar causas, utilize outro serviço contratual.')
    before=c.execute('SELECT * FROM rules WHERE tss=? AND cause=? AND form=? AND effective=?',(tss,cause,form,effective)).fetchone()
    key='tss_price:'+digest(tss)+':'+effective
    previous=c.execute('SELECT value FROM metadata WHERE key=?',(key,)).fetchone()
    c.execute('INSERT OR REPLACE INTO rules(tss,cause,form,effective,code,decision,evidence) VALUES(?,?,?,?,?,?,?)',(tss,cause,form,effective,code if decision=='Cobrar' else None,decision,evidence))
    if not cause and not form:
        # JSON null encerra um preço específico anterior e volta a seguir o preço da categoria.
        c.execute('INSERT OR REPLACE INTO metadata VALUES(?,?)',(key,json.dumps({'tss':tss,'effective':effective,'cents':amount},ensure_ascii=False)))
    return {'rule':dict(before) if before else None,'tss_price':json.loads(previous[0]) if previous else None}


def _source_stamp(root):
    source=Path(root).resolve()/'vistorias.sqlite3'
    if not source.exists():return [0,0]
    with closing(sqlite3.connect(source.as_uri()+'?mode=ro',uri=True)) as c:
        # Imports são append-only; não varrer visits para validar o formulário.
        return list(c.execute('SELECT COUNT(*),COALESCE(MAX(id),0) FROM imports').fetchone())


def _impact(c,root,change):
    month_bounds(change.get('effective',''));future(c,change['effective'])
    months={change['effective']}
    months.update(r[0] for r in c.execute('SELECT DISTINCT month FROM commercial WHERE month IS NOT NULL AND month NOT IN (SELECT month FROM approvals)'))
    months.update(r[0] for r in c.execute('SELECT month FROM geo WHERE month>=?',(change['effective'],)))
    def totals(month):
        p=preview(root,month,c)
        return {'inspection_cents':p['inspection_cents'],'geo_cents':p['geo']['cents'],'pending':p['pending'],'services':p['services'],'rates':p['rates'],'deadline':inspections.summary(root,*month_bounds(month),policy=p['policy'])[0]}
    before={m:totals(m) for m in sorted(months)}
    c.execute('SAVEPOINT proposed_catalog')
    try:
        _write_change(c,change);after={m:totals(m) for m in sorted(months)}
    finally:c.execute('ROLLBACK TO proposed_catalog');c.execute('RELEASE proposed_catalog')
    return [{'month':m,'before':before[m],'after':after[m],'delta_cents':after[m]['inspection_cents']+after[m]['geo_cents']-before[m]['inspection_cents']-before[m]['geo_cents']} for m in sorted(months) if m>=change['effective'] or before[m]!=after[m]]


def prepare_change(root,change):
    """Simulação exata em savepoint. Não aprova mês nem persiste regra/preço."""
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE')
        try:
            result={'root':str(Path(root).resolve()),'change':change,'revision':c.execute('SELECT value FROM revision WHERE id=1').fetchone()[0],'source':_source_stamp(root),'impact':_impact(c,root,change)}
            result['token']=digest(result);return result
        finally:c.rollback()


def confirm_change(root,plan,confirmed=False):
    if not confirmed:raise ValueError('Confirmação explícita do impacto obrigatória.')
    if plan.get('root')!=str(Path(root).resolve()) or plan.get('token')!=digest({k:v for k,v in plan.items() if k!='token'}):raise ValueError('Prévia inválida; gere novamente.')
    with connect(root) as c:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT value FROM revision WHERE id=1').fetchone()[0]!=plan['revision'] or _source_stamp(root)!=plan['source']:raise ValueError('Dados alterados; gere uma nova prévia de impacto.')
        actual=_impact(c,root,plan['change'])
        if actual!=plan['impact']:raise ValueError('Impacto mudou; confira novamente antes de confirmar.')
        before=_write_change(c,plan['change'])
        audit(c,'Alteração de catálogo confirmada',{'change':plan['change'],'before':before,'impact':actual,'token':plan['token']})
    return actual


def execution_stamp(value,day):
    stamp=inspections.date_value(value)
    if stamp is None:
        try:stamp=datetime.fromisoformat(str(value)).isoformat()
        except (ValueError,TypeError):return None
    return stamp if day and stamp[:10]==day else None


def _migrate_contract(c,root):
    """Backup verificado antes de alterar somente a projeção financeira."""
    if c.execute('PRAGMA user_version').fetchone()[0]>=2:return
    backup(Path(root).resolve()/'financeiro.sqlite3')
    try:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('PRAGMA user_version').fetchone()[0]>=2:c.rollback();return
        c.execute('ALTER TABLE commercial ADD COLUMN city_key TEXT')
        c.execute('ALTER TABLE commercial ADD COLUMN executed_at TEXT')
        c.execute('ALTER TABLE commercial ADD COLUMN month TEXT')
        cursor=c.execute('SELECT id,city,day,raw_hash FROM commercial ORDER BY id')
        # O timestamp vem do registro bruto preservado no banco operacional.
        source=Path(root).resolve()/'vistorias.sqlite3'
        if not source.exists() and c.execute('SELECT 1 FROM commercial LIMIT 1').fetchone():raise ValueError('Banco operacional ausente; migração cancelada sem alterar os dados.')
        if source.exists():
            with closing(sqlite3.connect(source.as_uri()+'?mode=ro',uri=True)) as op:
                while batch:=cursor.fetchmany(2000):
                    ids=[r['id'] for r in batch];raw={id:json.loads(text) for id,text in op.execute('SELECT id,raw FROM visits WHERE id IN ('+','.join('?' for _ in ids)+')',ids)}
                    if len(raw)!=len(ids):raise ValueError('Projeção sem origem correspondente; migração cancelada.')
                    c.executemany('UPDATE commercial SET city_key=?,executed_at=?,month=? WHERE id=?',[(norm(r['city']),execution_stamp(raw.get(r['id'],{}).get('Data Fim Execução'),r['day']),r['day'][:7] if r['day'] else None,r['id']) for r in batch])
        c.execute('CREATE INDEX commercial_month ON commercial(month,closed,city_key)')
        c.execute('CREATE INDEX commercial_scope ON commercial(city_key,month)')
        c.execute('CREATE INDEX commercial_month_keys ON commercial(month,contract,os)')
        c.execute('''CREATE TABLE frozen_charges(month TEXT NOT NULL,visit_id INTEGER NOT NULL,contract TEXT NOT NULL,code TEXT NOT NULL,os TEXT NOT NULL,cents INTEGER NOT NULL,PRIMARY KEY(month,visit_id))''')
        c.execute('CREATE INDEX frozen_os ON frozen_charges(contract,code,os,month)')
        for old in c.execute('SELECT month,payload FROM approvals').fetchall():
            sql,args=classified_sql(old['month'],c)
            calculated=c.execute("SELECT COALESCE(SUM(cents),0) FROM ("+sql+") WHERE reason='Faturável'",args).fetchone()[0]
            if calculated!=json.loads(old['payload'])['inspection_cents']:raise ValueError('Reconstrução histórica não coincide; migração cancelada, restaure/verifique o backup.')
            c.execute("INSERT INTO frozen_charges SELECT month,id,contract,code,os,cents FROM ("+sql+") WHERE reason='Faturável'",args)
        initial=contract_policy.initial('2026-09')
        c.execute('INSERT OR IGNORE INTO metadata VALUES(?,?)',('contract_scope:2026-09',json.dumps({'effective':'2026-09','allowed':initial['allowed'],'excluded':initial['excluded']},ensure_ascii=False)))
        for code,mode in initial['modes'].items():c.execute('INSERT OR IGNORE INTO metadata VALUES(?,?)',('billing_mode:'+code+':2026-09',json.dumps({'effective':'2026-09','code':code,'mode':mode,'representative':'latest'})))
        c.execute('PRAGMA user_version=2')
        audit(c,'Preparar escopo e cobrança contratual',{'version':2,'initial_scope':contract_policy.initial('2026-09'),'representative':'Última execução elegível; empate pelo menor ID original','historical_approvals_preserved':True})
        c.commit()
    except Exception:c.rollback();raise


def municipal_summary(root,month):
    return preview(root,month)['municipalities']


def export_period_catalog(root,month):
    main=export_catalog(root,month)
    main['period_catalogs']={}
    for m in main['contexts']:
        period=export_catalog(root,m,include_frozen=False);period.pop('frozen_charges',None);period.pop('contexts',None);main['period_catalogs'][m]=period
    source=Path(root).resolve()/'vistorias.sqlite3'
    if source.exists():
        with closing(sqlite3.connect(source.as_uri()+'?mode=ro',uri=True)) as c:main['import_order']=[r[0] for r in c.execute('SELECT sha FROM imports ORDER BY id')]
    return main
