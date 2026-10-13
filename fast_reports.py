"""Relatórios agregados rápidos e exportação completa em memória limitada."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import openpyxl
from openpyxl.cell import WriteOnlyCell
from core import CATEGORIES

EXECUTOR=ThreadPoolExecutor(max_workers=1)

def text_row(sheet, values):
    cells=[]
    for value in values:
        cell=WriteOnlyCell(sheet,value=value)
        if isinstance(value,str):cell.data_type='s'
        cells.append(cell)
    sheet.append(cells)

def workbook(month,summary,products,imports,audit,contract,closed,goals):
    w=openpyxl.Workbook(write_only=True);s=w.create_sheet('Fechamento')
    for row in [['Contrato',contract.get('name','Não informado')],['Número',contract.get('number','Não informado')],
                ['Referência',month],['Situação','Fechado' if closed else 'Prévia'],
                ['Código','Categoria','Quantidade','Valor unitário (R$)','Total (R$)']]:text_row(s,row)
    for r in summary:text_row(s,[r['Código'],r['Categoria'],r['Quantidade'],r['unit_cents']/100,r['total_cents']/100])
    text_row(s,['','TOTAL','','',sum(r['total_cents'] for r in summary)/100])
    for by,title in [('employee','Funcionários'),('city','Unidades'),('group','Grupos')]:
        s=w.create_sheet(title);text_row(s,[title,'Registros','Leituras faturáveis','Comunicados','Apontamentos','Total (R$)','Meta mensal','Atingimento (%)'])
        for r in products[by]:
            name=next(iter(r.values()));goal=goals.get(by+'|'+name)
            text_row(s,[*list(r.values())[:-1],r['total_cents']/100,goal,round(100*r['Leituras faturáveis']/goal,1) if goal else None])
    s=w.create_sheet('Arquivos e integridade');text_row(s,['Arquivo','Aba','SHA-256','Importado em UTC','Mapeamento','Situação'])
    for r in imports:text_row(s,[r['name'],r['sheet'],r['sha'],r['created_at'],r['mapping'],'Excluído' if r['deleted'] else 'Ativo'])
    s=w.create_sheet('Histórico');text_row(s,['Ação','Detalhe','Data UTC'])
    for r in audit:text_row(s,[r['action'],r['detail'],r['created_at']])
    return w

def summary_excel(*args):
    w=workbook(*args);out=BytesIO();w.save(out);return out.getvalue()

def detailed_excel(db_path,revision,month,prices,report_args):
    # A snapshot keeps the report consistent without locking live writes for the export.
    target=Path(db_path).parent/'exports';target.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=target) as folder:
        snapshot=Path(folder)/'snapshot.sqlite3'
        with closing(sqlite3.connect(db_path)) as src,closing(sqlite3.connect(snapshot)) as dst:src.backup(dst)
        with closing(sqlite3.connect(snapshot)) as c:
            if c.execute('SELECT revision FROM cache_revision').fetchone()[0]!=revision:
                raise ValueError('A base mudou antes da geração. Gere o relatório novamente.')
            c.row_factory=sqlite3.Row
            w=workbook(*report_args)
            raw=w.create_sheet('Dados importados');conference=w.create_sheet('Conferência')
            headers=list(dict.fromkeys(h for row in c.execute('SELECT headers FROM imports WHERE month=? ORDER BY id',(month,)) for h in json.loads(row['headers'])))
            text_row(raw,['ID do registro',*headers])
            text_row(conference,['ID','Arquivo','Aba','Linha original','Excluído','Motivo','Erros','Avisos',*CATEGORIES,'Total (R$)'])
            for r in c.execute('SELECT r.*,i.name,i.sheet FROM records r JOIN imports i ON i.id=r.import_id WHERE r.month=? ORDER BY r.id',(month,)):
                original=json.loads(r['raw']);result=json.loads(r['result']);u=result['units'];excluded=r['excluded'] or r['import_deleted']
                text_row(raw,[r['id'],*[original.get(h) for h in headers]])
                text_row(conference,[r['id'],r['name'],r['sheet'],r['line'],bool(excluded),r['reason'],'; '.join(result['errors']),'; '.join(result['alerts']),*[u[k] for k in CATEGORIES],0 if excluded else sum(u[k]*prices[k] for k in CATEGORIES)/100])
            out=target/f'fechamento-{month}-{revision}.xlsx';w.save(out)
    return out
