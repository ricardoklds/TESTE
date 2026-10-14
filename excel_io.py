"""Prévia pequena e leitura incremental. Não altera os arquivos originais."""
from io import BytesIO
import openpyxl
from core import FIELDS,norm,serial

def inspect(content):
    book=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=False)
    result=[]
    try:
        for sheet in book:
            sheet.reset_dimensions();headers=None;records=[]
            for number,values in enumerate(sheet.iter_rows(values_only=True),1):
                if headers is None:
                    names=[str(x).strip() if x is not None else '' for x in values]
                    if sum(norm(v) in {norm(x) for x in FIELDS.values()} for v in names)>=3:
                        if len(set(names))!=len(names) or any(not x for x in names):raise ValueError('Cabeçalhos vazios ou repetidos.')
                        headers=names;header_row=number
                    elif number>=40:break
                elif any(v not in (None,'') for v in values):
                    records.append({'line':number,'raw':{h:serial(values[i]) if i<len(values) else None for i,h in enumerate(headers)}})
                    if len(records)>=15:break
            if headers:
                result.append({'sheet':sheet.title,'headers':headers,'header_row':header_row,'records':records,
                               'mapping':{k:next((h for h in headers if norm(h)==norm(v)),None) for k,v in FIELDS.items()}})
                if not result[-1]['mapping']['city']:
                    candidates=[h for h in headers if norm(h) in ('cidade/atc','cidade / atc','cidade','municipio')]
                    if len(candidates)==1:result[-1]['mapping']['city']=candidates[0]
    finally:book.close()
    if not result:raise ValueError('Nenhum cabeçalho reconhecido.')
    return result

def iter_records(content,parsed):
    book=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=False)
    try:
        sheet=book[parsed['sheet']];sheet.reset_dimensions();headers=parsed['headers']
        for number,values in enumerate(sheet.iter_rows(values_only=True),1):
            if number<=parsed['header_row'] or not any(v not in (None,'') for v in values):continue
            if any(v is not None for v in values[len(headers):]):raise ValueError(f'Linha {number}: dados fora do cabeçalho.')
            yield {'line':number,'raw':{h:serial(values[i]) if i<len(values) else None for i,h in enumerate(headers)}}
    finally:book.close()
